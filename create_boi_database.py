import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

SRC = r"D:\claude\Project\raw\BOI 15 upgrade up Mar 20'26_finance confirm.xlsx"
OUT = r"D:\claude\Project\Output\BOI15_Upgrade_Database.xlsx"

wb_src = openpyxl.load_workbook(SRC, data_only=True)
ws = wb_src['Upgrade BOI10&11 to 15']

# Parse rows into flat records
records = []
current_process = ""
current_dept = ""

for r in range(2, ws.max_row + 1):
    a_val = ws.cell(r, 1).value
    b_val = ws.cell(r, 2).value

    # Detect section headers
    if a_val and isinstance(a_val, str):
        if 'ขั้นตอน ASSEMBLY' in a_val:
            current_dept = "ASSEMBLY"
            continue
        elif 'ขั้นตอน TESTING' in a_val:
            current_dept = "TESTING"
            continue
        elif a_val == 'ขั้นตอน':  # header row
            continue
        elif a_val == 'รวมทั้งสิ้น':
            continue

    # Skip rows with no upgrade detail
    if not b_val:
        continue

    # Update process name if column A has value
    if a_val and isinstance(a_val, str) and a_val not in ('ขั้นตอน',):
        current_process = a_val.strip()

    # Parse BOI column
    boi_raw = ws.cell(r, 15).value  # column O
    boi_str = str(boi_raw) if boi_raw else ""
    boi10 = "Yes" if "10" in boi_str else "No"
    boi11 = "Yes" if "11" in boi_str else "No"

    # Parse financial columns
    pr_val = ws.cell(r, 17).value   # Q
    spent_val = ws.cell(r, 18).value  # R
    remain_val = ws.cell(r, 19).value  # S
    total_usd = ws.cell(r, 12).value  # L

    # Determine status
    spent_num = float(spent_val) if spent_val else 0
    remain_num = float(remain_val) if remain_val else 0
    has_pr = pr_val is not None and str(pr_val).strip() != ""

    if spent_num > 0 and remain_num <= 0:
        status = "Completed"
    elif spent_num > 0 and remain_num > 0:
        status = "In Progress"
    elif has_pr and spent_num == 0:
        status = "PR Issued"
    else:
        status = "Pending"

    records.append({
        'dept': current_dept or "ASSEMBLY",
        'process': current_process,
        'upgrade_details': b_val,
        'propose': ws.cell(r, 3).value,
        'thai_desc': ws.cell(r, 4).value,
        'efficiency': ws.cell(r, 5).value,
        'mc_no': ws.cell(r, 6).value,
        'qty_set': ws.cell(r, 7).value,
        'qty_machine': ws.cell(r, 8).value,
        'cost_ea_baht': ws.cell(r, 9).value,
        'cost_ea_usd': ws.cell(r, 10).value,
        'total_baht': ws.cell(r, 11).value,
        'total_usd': total_usd,
        'contact': ws.cell(r, 13).value,
        'plant': ws.cell(r, 14).value,
        'boi10': boi10,
        'boi11': boi11,
        'ccid': ws.cell(r, 16).value,
        'pr': pr_val,
        'spent_usd': spent_val,
        'remain_usd': remain_val,
        'status': status,
        'remark': ws.cell(r, 20).value,
    })

print(f"Parsed {len(records)} records")

# Create output workbook
wb = Workbook()
ws_data = wb.active
ws_data.title = "BOI Upgrade Database"

# Styles
DARK_BLUE = "1B3A5C"
header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor=DARK_BLUE)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0')
)
alt_fill = PatternFill("solid", fgColor="E8F0FE")
white_fill = PatternFill("solid", fgColor="FFFFFF")

headers = [
    "No.", "Department", "Process", "Upgrade Details", "Propose (EN)",
    "รายละเอียด (TH)", "ประสิทธิภาพ", "MC No.", "Qty/Set", "Qty Machine",
    "Cost/ea (Baht)", "Cost/ea (USD)", "Total (Baht)", "Total (USD)",
    "Contact", "Plant", "BOI 10", "BOI 11", "CCID", "PR#",
    "Spent (USD)", "Remain (USD)", "% Spent", "Status", "Remark"
]

# Write headers
for c, h in enumerate(headers, 1):
    cell = ws_data.cell(1, c, h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write data
for i, rec in enumerate(records):
    row = i + 2
    data_font = Font(name="Arial", size=10)

    ws_data.cell(row, 1, i + 1)  # No.
    ws_data.cell(row, 2, rec['dept'])
    ws_data.cell(row, 3, rec['process'])
    ws_data.cell(row, 4, rec['upgrade_details'])
    ws_data.cell(row, 5, rec['propose'])
    ws_data.cell(row, 6, rec['thai_desc'])
    ws_data.cell(row, 7, rec['efficiency'])
    ws_data.cell(row, 8, rec['mc_no'])
    ws_data.cell(row, 9, rec['qty_set'])
    ws_data.cell(row, 10, rec['qty_machine'])
    ws_data.cell(row, 11, rec['cost_ea_baht'])
    ws_data.cell(row, 12, rec['cost_ea_usd'])
    ws_data.cell(row, 13, rec['total_baht'])
    ws_data.cell(row, 14, rec['total_usd'])
    ws_data.cell(row, 15, rec['contact'])
    ws_data.cell(row, 16, rec['plant'])
    ws_data.cell(row, 17, rec['boi10'])
    ws_data.cell(row, 18, rec['boi11'])
    ws_data.cell(row, 19, rec['ccid'])
    ws_data.cell(row, 20, rec['pr'])
    ws_data.cell(row, 21, rec['spent_usd'])
    ws_data.cell(row, 22, rec['remain_usd'])
    # % Spent formula
    ws_data.cell(row, 23).value = f'=IF(N{row}=0,"",U{row}/N{row})'
    ws_data.cell(row, 24, rec['status'])
    ws_data.cell(row, 25, rec['remark'])

    # Apply formatting
    fill = alt_fill if i % 2 == 0 else white_fill
    for c in range(1, 26):
        cell = ws_data.cell(row, c)
        cell.font = data_font
        cell.border = thin_border
        cell.fill = fill
        if c in (1, 9, 10):
            cell.alignment = Alignment(horizontal="center")
        elif c in (2, 3, 15, 16, 17, 18, 24):
            cell.alignment = Alignment(horizontal="center")

# Number formats
last_row = len(records) + 1
for r in range(2, last_row + 1):
    ws_data.cell(r, 11).number_format = '#,##0'       # Baht ea
    ws_data.cell(r, 12).number_format = '#,##0.00'     # USD ea
    ws_data.cell(r, 13).number_format = '#,##0'        # Total Baht
    ws_data.cell(r, 14).number_format = '#,##0.00'     # Total USD
    ws_data.cell(r, 21).number_format = '#,##0.00'     # Spent USD
    ws_data.cell(r, 22).number_format = '#,##0.00'     # Remain USD
    ws_data.cell(r, 23).number_format = '0.0%'         # % Spent

# Conditional formatting for Status (column X = 24)
green_fill = PatternFill("solid", fgColor="C6EFCE")
green_font = Font(name="Arial", size=10, color="006100")
yellow_fill = PatternFill("solid", fgColor="FFEB9C")
yellow_font = Font(name="Arial", size=10, color="9C6500")
orange_fill = PatternFill("solid", fgColor="FCD5B4")
orange_font = Font(name="Arial", size=10, color="974706")
gray_fill = PatternFill("solid", fgColor="E0E0E0")
gray_font = Font(name="Arial", size=10, color="666666")

status_range = f"X2:X{last_row}"
ws_data.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Completed"'], fill=green_fill, font=green_font))
ws_data.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"In Progress"'], fill=yellow_fill, font=yellow_font))
ws_data.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"PR Issued"'], fill=orange_fill, font=orange_font))
ws_data.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Pending"'], fill=gray_fill, font=gray_font))

# Column widths
col_widths = {
    1: 5, 2: 13, 3: 18, 4: 45, 5: 40, 6: 40, 7: 35,
    8: 30, 9: 9, 10: 12, 11: 14, 12: 13, 13: 14, 14: 13,
    15: 20, 16: 8, 17: 8, 18: 8, 19: 12, 20: 18,
    21: 13, 22: 13, 23: 10, 24: 13, 25: 35
}
for c, w in col_widths.items():
    ws_data.column_dimensions[get_column_letter(c)].width = w

# AutoFilter
ws_data.auto_filter.ref = f"A1:{get_column_letter(25)}{last_row}"

# Freeze top row
ws_data.freeze_panes = "A2"

# ========== SUMMARY SHEET ==========
ws_sum = wb.create_sheet("Summary")

sum_header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
sum_header_fill = PatternFill("solid", fgColor=DARK_BLUE)
sum_data_font = Font(name="Arial", size=10)
sum_bold = Font(name="Arial", size=10, bold=True)
accent_fill = PatternFill("solid", fgColor="2E75B6")

# Title
ws_sum.cell(1, 1, "BOI 15 Upgrade Summary by Department & Process")
ws_sum.cell(1, 1).font = Font(name="Arial", bold=True, size=14, color=DARK_BLUE)
ws_sum.merge_cells("A1:H1")

# Build summary data
from collections import defaultdict
summary = defaultdict(lambda: {'count': 0, 'total_usd': 0, 'spent_usd': 0, 'remain_usd': 0,
                                'completed': 0, 'in_progress': 0, 'pending': 0})
for rec in records:
    key = (rec['dept'], rec['process'])
    s = summary[key]
    s['count'] += 1
    s['total_usd'] += float(rec['total_usd'] or 0)
    s['spent_usd'] += float(rec['spent_usd'] or 0)
    s['remain_usd'] += float(rec['remain_usd'] or 0)
    if rec['status'] == 'Completed':
        s['completed'] += 1
    elif rec['status'] in ('In Progress', 'PR Issued'):
        s['in_progress'] += 1
    else:
        s['pending'] += 1

sum_headers = ["Department", "Process", "Items", "Total (USD)", "Spent (USD)", "Remain (USD)", "% Spent", "Status Breakdown"]
sr = 3
for c, h in enumerate(sum_headers, 1):
    cell = ws_sum.cell(sr, c, h)
    cell.font = sum_header_font
    cell.fill = sum_header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

sr = 4
sorted_keys = sorted(summary.keys(), key=lambda k: (0 if k[0] == 'ASSEMBLY' else 1, k[1]))

dept_totals = defaultdict(lambda: {'count': 0, 'total_usd': 0, 'spent_usd': 0, 'remain_usd': 0})
current_dept_sum = ""

for key in sorted_keys:
    dept, proc = key
    s = summary[key]

    # Insert department subtotal when department changes
    if current_dept_sum and dept != current_dept_sum:
        dt = dept_totals[current_dept_sum]
        ws_sum.cell(sr, 1, f"{current_dept_sum} Total")
        ws_sum.cell(sr, 3, dt['count'])
        ws_sum.cell(sr, 4, dt['total_usd'])
        ws_sum.cell(sr, 5, dt['spent_usd'])
        ws_sum.cell(sr, 6, dt['remain_usd'])
        ws_sum.cell(sr, 7).value = f'=IF(D{sr}=0,"",E{sr}/D{sr})'
        ws_sum.cell(sr, 7).number_format = '0.0%'
        for c in range(1, 9):
            ws_sum.cell(sr, c).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            ws_sum.cell(sr, c).fill = PatternFill("solid", fgColor="2E75B6")
            ws_sum.cell(sr, c).border = thin_border
        sr += 1

    current_dept_sum = dept
    dept_totals[dept]['count'] += s['count']
    dept_totals[dept]['total_usd'] += s['total_usd']
    dept_totals[dept]['spent_usd'] += s['spent_usd']
    dept_totals[dept]['remain_usd'] += s['remain_usd']

    ws_sum.cell(sr, 1, dept)
    ws_sum.cell(sr, 2, proc)
    ws_sum.cell(sr, 3, s['count'])
    ws_sum.cell(sr, 4, s['total_usd'])
    ws_sum.cell(sr, 5, s['spent_usd'])
    ws_sum.cell(sr, 6, s['remain_usd'])
    ws_sum.cell(sr, 7).value = f'=IF(D{sr}=0,"",E{sr}/D{sr})'
    ws_sum.cell(sr, 7).number_format = '0.0%'
    status_text = f"Done:{s['completed']} | WIP:{s['in_progress']} | Pending:{s['pending']}"
    ws_sum.cell(sr, 8, status_text)

    fill_row = alt_fill if (sr - 4) % 2 == 0 else white_fill
    for c in range(1, 9):
        cell = ws_sum.cell(sr, c)
        cell.font = sum_data_font
        cell.border = thin_border
        cell.fill = fill_row
        if c in (3,):
            cell.alignment = Alignment(horizontal="center")
        if c in (4, 5, 6):
            cell.number_format = '#,##0.00'

    sr += 1

# Last department subtotal
if current_dept_sum:
    dt = dept_totals[current_dept_sum]
    ws_sum.cell(sr, 1, f"{current_dept_sum} Total")
    ws_sum.cell(sr, 3, dt['count'])
    ws_sum.cell(sr, 4, dt['total_usd'])
    ws_sum.cell(sr, 5, dt['spent_usd'])
    ws_sum.cell(sr, 6, dt['remain_usd'])
    ws_sum.cell(sr, 7).value = f'=IF(D{sr}=0,"",E{sr}/D{sr})'
    ws_sum.cell(sr, 7).number_format = '0.0%'
    for c in range(1, 9):
        ws_sum.cell(sr, c).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        ws_sum.cell(sr, c).fill = PatternFill("solid", fgColor="2E75B6")
        ws_sum.cell(sr, c).border = thin_border
    sr += 1

# Grand total
sr += 1
grand = {'count': 0, 'total_usd': 0, 'spent_usd': 0, 'remain_usd': 0}
for dt in dept_totals.values():
    grand['count'] += dt['count']
    grand['total_usd'] += dt['total_usd']
    grand['spent_usd'] += dt['spent_usd']
    grand['remain_usd'] += dt['remain_usd']

ws_sum.cell(sr, 1, "GRAND TOTAL")
ws_sum.cell(sr, 3, grand['count'])
ws_sum.cell(sr, 4, grand['total_usd'])
ws_sum.cell(sr, 5, grand['spent_usd'])
ws_sum.cell(sr, 6, grand['remain_usd'])
ws_sum.cell(sr, 7).value = f'=IF(D{sr}=0,"",E{sr}/D{sr})'
ws_sum.cell(sr, 7).number_format = '0.0%'
for c in range(1, 9):
    ws_sum.cell(sr, c).font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws_sum.cell(sr, c).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws_sum.cell(sr, c).border = thin_border
    if c in (4, 5, 6):
        ws_sum.cell(sr, c).number_format = '#,##0.00'

# Summary column widths
sum_widths = {1: 18, 2: 22, 3: 8, 4: 15, 5: 15, 6: 15, 7: 10, 8: 30}
for c, w in sum_widths.items():
    ws_sum.column_dimensions[get_column_letter(c)].width = w

ws_sum.freeze_panes = "A4"

# Save
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Saved to {OUT}")
print(f"Total records: {len(records)}")
print(f"Departments: {set(r['dept'] for r in records)}")
print(f"Processes: {set(r['process'] for r in records)}")
statuses = defaultdict(int)
for r in records:
    statuses[r['status']] += 1
print(f"Status counts: {dict(statuses)}")
