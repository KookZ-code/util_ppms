# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
import lxml.etree as etree
import openpyxl
from collections import Counter

# Load data
wb = openpyxl.load_workbook('D:/claude/Project/raw/UPDATE STATUS MACHINE DIE ATTACHล่าสุด 7.xlsx', data_only=True)

# ---- Parse Update Apr 04 (AD889 Turn-on Tracking) ----
ws_update = wb["Update Apr 04'26 MachienDA "]
update_machines = []
for row in ws_update.iter_rows(min_row=2, max_row=25, values_only=False):
    item = row[0].value
    number = row[1].value
    model = row[2].value
    usable = row[3].value
    pkg = row[4].value
    status = row[5].value
    remark = row[6].value
    if item and isinstance(item, (int, float)) and number:
        update_machines.append({
            'number': number, 'model': model or '', 'usable': usable or '',
            'pkg': pkg or '', 'status': (status or '').strip(), 'remark': remark or ''
        })

# Update status counts
upd_status_clean = Counter()
for m in update_machines:
    s = m['status']
    if 'Run' in s:
        upd_status_clean['Running'] += 1
    elif 'Stand by' in s:
        upd_status_clean['Stand By'] += 1
    elif 'boot' in s.lower():
        upd_status_clean['Boot Fail'] += 1
    elif 'Wait' in s:
        upd_status_clean['Wait Setup/PC'] += 1
    elif 'On going' in s:
        upd_status_clean['On Going (Repair)'] += 1
    elif s:
        upd_status_clean[s] += 1
    else:
        upd_status_clean['Other'] += 1

upd_usable_yes = sum(1 for m in update_machines if m['usable'] == 'Yes')
upd_usable_no = sum(1 for m in update_machines if m['usable'] == 'No')

# ---- Parse Die Attach (Full Inventory) ----
ws_da = wb['Die Attach Machine']
da_machines = []
for row in ws_da.iter_rows(min_row=4, max_row=117, values_only=False):
    num = row[0].value
    model = row[1].value
    condition = row[2].value
    status = row[3].value
    remark = row[4].value
    if num:
        da_machines.append({
            'number': num, 'model': model, 'condition': condition,
            'status': status or '', 'remark': remark or ''
        })

# ---- Parse SAW ----
ws_saw = wb['SAW']
saw_machines = []
for row in ws_saw.iter_rows(min_row=4, max_row=56, values_only=False):
    num = row[0].value
    name = row[1].value
    model = row[2].value
    cond = row[3].value
    status = row[4].value
    remark = row[5].value
    if num:
        saw_machines.append({
            'number': num, 'name': name or '', 'model': (model or '').strip(),
            'condition': cond or '', 'status': status or '', 'remark': remark or ''
        })

# ---- Statistics ----
da_total = len(da_machines)
da_normal = sum(1 for m in da_machines if m['condition'] == 'ปกติ')
da_broken = sum(1 for m in da_machines if m['condition'] == 'เสีย')

da_status_clean = Counter()
for m in da_machines:
    s = m['status'].strip() if m['status'] else ''
    if not s:
        if 'Next plan' in str(m.get('remark', '')):
            s = 'Next Plan'
        elif 'Wait' in str(m.get('remark', '')):
            s = 'Wait Scrap'
    if 'Run' in s:
        da_status_clean['Running'] += 1
    elif 'Boot' in s:
        da_status_clean['Boot Fail'] += 1
    elif 'Stand' in s:
        da_status_clean['On Stand By'] += 1
    elif 'Scrap' in s or 'Scarp' in s:
        da_status_clean['Wait Scrap'] += 1
    elif 'Next' in s or s == '':
        da_status_clean['Next Plan / Other'] += 1
    else:
        da_status_clean[s if s else 'Other'] += 1

da_model_count = Counter(m['model'] for m in da_machines)
da_model_normal = Counter()
da_model_broken = Counter()
for m in da_machines:
    if m['condition'] == 'ปกติ':
        da_model_normal[m['model']] += 1
    elif m['condition'] == 'เสีย':
        da_model_broken[m['model']] += 1

saw_total = len(saw_machines)
saw_normal = sum(1 for m in saw_machines if m['condition'] in ('ปกติ', 'ได้ dummy tape ok'))
saw_broken_or_issue = saw_total - saw_normal
saw_model_count = Counter(m['model'] for m in saw_machines)

# Colors
DARK_BLUE = RGBColor(0x0E, 0x36, 0x89)
MED_BLUE = RGBColor(0x1D, 0x9C, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x5E, 0xBF, 0x33)
RED = RGBColor(0xCC, 0x00, 0x00)
ORANGE = RGBColor(0xFD, 0x7F, 0x20)
GRAY = RGBColor(0x8A, 0x8A, 0x8A)
BLACK = RGBColor(0x0A, 0x0B, 0x0F)
LIGHT_GRAY = RGBColor(0xD9, 0xD9, 0xD9)
BG_LIGHT = RGBColor(0xF7, 0xF7, 0xF7)
PURPLE = RGBColor(0x70, 0x20, 0x76)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=BLACK, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_shape_box(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def add_kpi_card(slide, left, top, width, height, title, value, color):
    add_shape_box(slide, left, top, width, height, WHITE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    add_text_box(slide, left+0.15, top+0.15, width-0.3, 0.7, str(value), font_size=34, bold=True, color=color, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, left+0.15, top+0.9, width-0.3, 0.5, title, font_size=13, bold=False, color=GRAY, alignment=PP_ALIGN.CENTER)

def add_header_bar(slide, title):
    hbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85))
    hbar.fill.solid()
    hbar.fill.fore_color.rgb = DARK_BLUE
    hbar.line.fill.background()
    add_text_box(slide, 0.5, 0.13, 12, 0.6, title, font_size=26, bold=True, color=WHITE)

def add_footer(slide, slide_num):
    add_text_box(slide, 0.3, 7.1, 1, 0.3, str(slide_num), font_size=9, color=GRAY, alignment=PP_ALIGN.LEFT)
    add_text_box(slide, 4.5, 7.1, 4.5, 0.3, 'Microchip Proprietary and Confidential', font_size=9, color=GRAY, alignment=PP_ALIGN.CENTER)

def set_chart_series_color(series_elem, hex_color):
    spPr = etree.SubElement(series_elem, qn('c:spPr'))
    sf = etree.SubElement(spPr, qn('a:solidFill'))
    sc = etree.SubElement(sf, qn('a:srgbClr'))
    sc.set('val', hex_color)

def set_chart_point_color(series_elem, idx, hex_color):
    dPt = etree.SubElement(series_elem, qn('c:dPt'))
    idx_elem = etree.SubElement(dPt, qn('c:idx'))
    idx_elem.set('val', str(idx))
    spPr = etree.SubElement(dPt, qn('c:spPr'))
    solidFill = etree.SubElement(spPr, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', hex_color)

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ========== SLIDE 1: Title ==========
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide1, WHITE)
# Top accent bar
top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
top_bar.fill.solid()
top_bar.fill.fore_color.rgb = DARK_BLUE
top_bar.line.fill.background()
# Left accent block
left_block = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.15), Inches(0.4), Inches(7.35))
left_block.fill.solid()
left_block.fill.fore_color.rgb = DARK_BLUE
left_block.line.fill.background()
# Title area
add_shape_box(slide1, 1.5, 1.8, 10.5, 3.6, BG_LIGHT)
add_text_box(slide1, 2.0, 2.1, 9.5, 1.0, 'ASSY Machine Status Report', font_size=42, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.LEFT)
add_text_box(slide1, 2.0, 3.0, 9.5, 0.7, 'Die Attach & SAW Machine Capacity Summary', font_size=22, bold=False, color=MED_BLUE, alignment=PP_ALIGN.LEFT)
# Divider line
div_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.0), Inches(3.7), Inches(3), Inches(0.04))
div_line.fill.solid()
div_line.fill.fore_color.rgb = MED_BLUE
div_line.line.fill.background()
add_text_box(slide1, 2.0, 4.0, 9.5, 0.5, 'Update: April 4, 2026', font_size=18, bold=False, color=GRAY, alignment=PP_ALIGN.LEFT)
add_text_box(slide1, 2.0, 4.5, 9.5, 0.5, 'MTAI - Assembly & Cap Engineering', font_size=14, bold=False, color=GRAY, alignment=PP_ALIGN.LEFT)
add_footer(slide1, 1)

# ========== SLIDE 2: Executive Summary ==========
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2, BG_LIGHT)
add_header_bar(slide2, 'Executive Summary')

total_all = da_total + saw_total
add_kpi_card(slide2, 0.5, 1.2, 2.4, 1.4, 'Total Machines', total_all, MED_BLUE)
add_kpi_card(slide2, 3.2, 1.2, 2.4, 1.4, 'Die Attach', da_total, DARK_BLUE)
add_kpi_card(slide2, 5.9, 1.2, 2.4, 1.4, 'SAW Machines', saw_total, DARK_BLUE)
add_kpi_card(slide2, 8.6, 1.2, 2.4, 1.4, 'DA Normal', da_normal, GREEN)
add_kpi_card(slide2, 11.3, 1.2, 2.0, 1.4, 'DA Issues', da_broken, RED)

# Summary box
add_shape_box(slide2, 0.5, 2.9, 12.3, 4.2, WHITE)
bullets = [
    f'Die Attach: {da_total} machines total - {da_normal} Normal (ปกติ), {da_broken} Broken (เสีย)',
    f'Die Attach Status: {da_status_clean.get("Running",0)} Running, {da_status_clean.get("Boot Fail",0)} Boot Fail, {da_status_clean.get("On Stand By",0)} On Stand By, {da_status_clean.get("Wait Scrap",0)} Wait Scrap',
    f'AD889 Turn-on Progress (Apr 04): {upd_status_clean.get("Running",0)} Running, {upd_status_clean.get("Boot Fail",0)} Boot Fail, {upd_status_clean.get("Stand By",0)} Stand By, {upd_status_clean.get("Wait Setup/PC",0)} Wait Setup/PC',
    f'AD889 Usable: {upd_usable_yes} Yes, {upd_usable_no} No out of {len(update_machines)} machines tracked',
    f'SAW: {saw_total} machines total - {saw_normal} Operational, {saw_broken_or_issue} with Issues',
    f'Key SAW Issues: Vacuum fail (board discontinued), Spindle water leak, HDD fail (obsolete SAS), CO2 Bubbler fail',
    f'Die Attach Models: {len(da_model_count)} types - {", ".join(sorted(da_model_count.keys())[:6])}...',
    f'SAW Models: {", ".join(sorted(saw_model_count.keys()))}',
]
for i, bullet in enumerate(bullets):
    add_text_box(slide2, 0.9, 3.1 + i*0.48, 11.3, 0.42, '\u2022  ' + bullet, font_size=13, bold=False, color=BLACK)

add_footer(slide2, 2)

# ========== SLIDE 3: AD889 Turn-on Status (Apr 04) ==========
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3, BG_LIGHT)
add_header_bar(slide3, "AD889 Turn-on Status - Update Apr 04'26")

# KPI cards for turn-on
add_kpi_card(slide3, 0.5, 1.1, 2.0, 1.3, 'Total Tracked', len(update_machines), MED_BLUE)
add_kpi_card(slide3, 2.7, 1.1, 2.0, 1.3, 'Running', upd_status_clean.get('Running', 0), GREEN)
add_kpi_card(slide3, 4.9, 1.1, 2.0, 1.3, 'Stand By', upd_status_clean.get('Stand By', 0), ORANGE)
add_kpi_card(slide3, 7.1, 1.1, 2.0, 1.3, 'Boot Fail', upd_status_clean.get('Boot Fail', 0), RED)
add_kpi_card(slide3, 9.3, 1.1, 2.0, 1.3, 'Wait Setup/PC', upd_status_clean.get('Wait Setup/PC', 0), PURPLE)
add_kpi_card(slide3, 11.5, 1.1, 1.6, 1.3, 'On Going', upd_status_clean.get('On Going (Repair)', 0), ORANGE)

# Status bar chart
upd_chart_data = CategoryChartData()
upd_labels = list(upd_status_clean.keys())
upd_values = list(upd_status_clean.values())
upd_chart_data.categories = upd_labels
upd_chart_data.add_series('Count', upd_values)
upd_chart = slide3.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2.7), Inches(6), Inches(3.8), upd_chart_data
).chart
upd_chart.has_legend = False
upd_plot = upd_chart.plots[0]
upd_plot.has_data_labels = True
upd_plot.data_labels.number_format = '0'
upd_plot.data_labels.font.size = Pt(11)
upd_chart.category_axis.tick_labels.font.size = Pt(10)
upd_chart.value_axis.visible = False
# Color each bar
upd_colors = {'Running': '5EBF33', 'Stand By': 'FD7F20', 'Boot Fail': 'CC0000',
              'Wait Setup/PC': '702076', 'On Going (Repair)': 'FD7F20', 'Other': '8A8A8A'}
for i, label in enumerate(upd_labels):
    set_chart_point_color(upd_chart.series[0]._element, i, upd_colors.get(label, '1D9CE4'))

add_text_box(slide3, 0.5, 6.6, 6, 0.4, 'AD889 Machine Status Distribution', font_size=13, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)

# Machine detail table
add_text_box(slide3, 6.8, 2.7, 6.2, 0.35, 'Machine Detail', font_size=15, bold=True, color=DARK_BLUE)
# Header
y = 3.15
hdr = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(y), Inches(6.2), Inches(0.32))
hdr.fill.solid()
hdr.fill.fore_color.rgb = DARK_BLUE
hdr.line.fill.background()
add_text_box(slide3, 6.85, y, 1.0, 0.32, 'Machine', font_size=10, bold=True, color=WHITE)
add_text_box(slide3, 7.85, y, 0.9, 0.32, 'Usable', font_size=10, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide3, 8.75, y, 1.1, 0.32, 'Pkg', font_size=10, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide3, 9.85, y, 1.2, 0.32, 'Status', font_size=10, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide3, 11.05, y, 1.95, 0.32, 'Remark', font_size=10, bold=True, color=WHITE)
y += 0.34
for i, m in enumerate(update_machines[:12]):
    row_color = BG_LIGHT if i % 2 == 0 else WHITE
    rbg = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(y), Inches(6.2), Inches(0.27))
    rbg.fill.solid()
    rbg.fill.fore_color.rgb = row_color
    rbg.line.fill.background()
    slide3.shapes._spTree.remove(rbg._element)
    slide3.shapes._spTree.insert(2, rbg._element)

    s_color = GREEN if 'Run' in m['status'] else RED if 'Boot' in m['status'].lower() or 'fail' in m['status'].lower() else ORANGE if 'Stand' in m['status'] else PURPLE if 'Wait' in m['status'] else BLACK
    add_text_box(slide3, 6.85, y, 1.0, 0.27, m['number'], font_size=9, color=BLACK)
    add_text_box(slide3, 7.85, y, 0.9, 0.27, m['usable'], font_size=9, color=GREEN if m['usable']=='Yes' else RED, alignment=PP_ALIGN.CENTER)
    add_text_box(slide3, 8.75, y, 1.1, 0.27, m['pkg'], font_size=9, color=BLACK, alignment=PP_ALIGN.CENTER)
    add_text_box(slide3, 9.85, y, 1.2, 0.27, m['status'], font_size=9, bold=True, color=s_color, alignment=PP_ALIGN.CENTER)
    remark_short = m['remark'][:35] + '...' if len(m['remark']) > 35 else m['remark']
    add_text_box(slide3, 11.05, y, 1.95, 0.27, remark_short, font_size=8, color=GRAY)
    y += 0.29

add_footer(slide3, 3)

# ========== SLIDE 4: Die Attach Overview ==========
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4, BG_LIGHT)
add_header_bar(slide4, 'Die Attach Machine - Full Inventory Overview')

# Condition pie chart
chart_data1 = CategoryChartData()
chart_data1.categories = ['Normal (ปกติ)', 'Broken (เสีย)']
chart_data1.add_series('Condition', (da_normal, da_broken))
chart1 = slide4.shapes.add_chart(
    XL_CHART_TYPE.PIE, Inches(0.5), Inches(1.2), Inches(5.5), Inches(3.8), chart_data1
).chart
chart1.has_legend = True
chart1.legend.position = XL_LEGEND_POSITION.BOTTOM
chart1.legend.include_in_layout = False
plot1 = chart1.plots[0]
plot1.has_data_labels = True
plot1.data_labels.number_format = '0'
plot1.data_labels.font.size = Pt(12)
set_chart_point_color(chart1.series[0]._element, 0, '5EBF33')
set_chart_point_color(chart1.series[0]._element, 1, 'CC0000')

add_text_box(slide4, 0.5, 5.1, 5.5, 0.4,
    f'Normal: {da_normal} ({da_normal*100//da_total}%)  |  Broken: {da_broken} ({da_broken*100//da_total}%)',
    font_size=13, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)

# Status bar chart
chart_data2 = CategoryChartData()
status_labels = list(da_status_clean.keys())
status_values = list(da_status_clean.values())
chart_data2.categories = status_labels
chart_data2.add_series('Count', status_values)
chart2 = slide4.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.5), Inches(1.2), Inches(6.3), Inches(3.8), chart_data2
).chart
chart2.has_legend = False
plot2 = chart2.plots[0]
plot2.has_data_labels = True
plot2.data_labels.number_format = '0'
plot2.data_labels.font.size = Pt(11)
chart2.category_axis.tick_labels.font.size = Pt(10)
chart2.value_axis.visible = False
# Color bars
da_bar_colors = {'Running': '5EBF33', 'Boot Fail': 'CC0000', 'On Stand By': 'FD7F20',
                 'Wait Scrap': '8A8A8A', 'Next Plan / Other': '1D9CE4'}
for i, label in enumerate(status_labels):
    set_chart_point_color(chart2.series[0]._element, i, da_bar_colors.get(label, '0E3689'))

add_text_box(slide4, 6.5, 5.1, 6.3, 0.4, 'Machine Status Distribution',
    font_size=13, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)

# Summary text
add_shape_box(slide4, 0.5, 5.7, 12.3, 1.2, WHITE)
add_text_box(slide4, 0.8, 5.8, 11.8, 0.4,
    f'Total Die Attach: {da_total}  |  Models: {", ".join(sorted(da_model_count.keys()))}',
    font_size=12, color=GRAY, alignment=PP_ALIGN.LEFT)
add_text_box(slide4, 0.8, 6.2, 11.8, 0.4,
    f'Running: {da_status_clean.get("Running",0)}  |  Boot Fail: {da_status_clean.get("Boot Fail",0)}  |  On Stand By: {da_status_clean.get("On Stand By",0)}  |  Next Plan: {da_status_clean.get("Next Plan / Other",0)}',
    font_size=12, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.LEFT)

add_footer(slide4, 4)

# ========== SLIDE 5: Die Attach by Model ==========
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5, BG_LIGHT)
add_header_bar(slide5, 'Die Attach Machine - By Model')

sorted_models = sorted(da_model_count.keys(), key=lambda x: da_model_count[x], reverse=True)
chart_data3 = CategoryChartData()
chart_data3.categories = sorted_models
chart_data3.add_series('Normal', [da_model_normal.get(m, 0) for m in sorted_models])
chart_data3.add_series('Broken', [da_model_broken.get(m, 0) for m in sorted_models])
chart3 = slide5.shapes.add_chart(
    XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.2), Inches(8), Inches(5.5), chart_data3
).chart
chart3.has_legend = True
chart3.legend.position = XL_LEGEND_POSITION.BOTTOM
plot3 = chart3.plots[0]
plot3.has_data_labels = True
plot3.data_labels.number_format = '0'
plot3.data_labels.font.size = Pt(10)
set_chart_series_color(chart3.series[0]._element, '5EBF33')
set_chart_series_color(chart3.series[1]._element, 'CC0000')

# Model summary table
add_text_box(slide5, 9.0, 1.1, 4, 0.35, 'Model Summary', font_size=15, bold=True, color=DARK_BLUE)
y = 1.55
hdr_bg = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(y), Inches(4.1), Inches(0.32))
hdr_bg.fill.solid()
hdr_bg.fill.fore_color.rgb = DARK_BLUE
hdr_bg.line.fill.background()
add_text_box(slide5, 9.05, y, 1.7, 0.32, '  Model', font_size=11, bold=True, color=WHITE)
add_text_box(slide5, 10.75, y, 0.7, 0.32, 'Total', font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide5, 11.45, y, 0.7, 0.32, 'OK', font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide5, 12.15, y, 0.7, 0.32, 'NG', font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
y += 0.34
for i, mdl in enumerate(sorted_models):
    total_m = da_model_count[mdl]
    ok_m = da_model_normal.get(mdl, 0)
    ng_m = da_model_broken.get(mdl, 0)
    row_color = BG_LIGHT if i % 2 == 0 else WHITE
    rbg = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(y), Inches(4.1), Inches(0.30))
    rbg.fill.solid()
    rbg.fill.fore_color.rgb = row_color
    rbg.line.fill.background()
    slide5.shapes._spTree.remove(rbg._element)
    slide5.shapes._spTree.insert(2, rbg._element)
    add_text_box(slide5, 9.05, y, 1.7, 0.30, '  ' + mdl, font_size=10, color=BLACK)
    add_text_box(slide5, 10.75, y, 0.7, 0.30, str(total_m), font_size=10, color=BLACK, alignment=PP_ALIGN.CENTER)
    add_text_box(slide5, 11.45, y, 0.7, 0.30, str(ok_m), font_size=10, color=GREEN, alignment=PP_ALIGN.CENTER)
    add_text_box(slide5, 12.15, y, 0.7, 0.30, str(ng_m), font_size=10, color=RED if ng_m > 0 else GRAY, alignment=PP_ALIGN.CENTER)
    y += 0.32

add_footer(slide5, 5)

# ========== SLIDE 6: SAW Overview ==========
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6, BG_LIGHT)
add_header_bar(slide6, 'SAW Machine - Overview')

# SAW model pie chart
chart_data4 = CategoryChartData()
saw_sorted_models = sorted(saw_model_count.keys(), key=lambda x: saw_model_count[x], reverse=True)
chart_data4.categories = saw_sorted_models
chart_data4.add_series('Count', [saw_model_count[m] for m in saw_sorted_models])
chart4 = slide6.shapes.add_chart(
    XL_CHART_TYPE.PIE, Inches(0.5), Inches(1.2), Inches(5.5), Inches(4.2), chart_data4
).chart
chart4.has_legend = True
chart4.legend.position = XL_LEGEND_POSITION.BOTTOM
chart4.legend.include_in_layout = False
plot4 = chart4.plots[0]
plot4.has_data_labels = True
plot4.data_labels.number_format = '0'
plot4.data_labels.font.size = Pt(11)
saw_pie_colors = ['0E3689', '1D9CE4', '5EBF33', 'FD7F20', 'CC0000', '702076', 'FFD53A']
for i in range(len(saw_sorted_models)):
    set_chart_point_color(chart4.series[0]._element, i, saw_pie_colors[i % len(saw_pie_colors)])

add_text_box(slide6, 0.5, 5.5, 5.5, 0.4, f'Total SAW Machines: {saw_total}',
    font_size=13, bold=True, color=DARK_BLUE, alignment=PP_ALIGN.CENTER)

# SAW KPI cards
saw_ok = sum(1 for m in saw_machines if m['condition'] == 'ปกติ')
saw_dummy_ok = sum(1 for m in saw_machines if 'dummy' in m['condition'].lower())
saw_shutdown = sum(1 for m in saw_machines if 'SHUTDOWN' in m['status'].upper())
saw_remark_issue = sum(1 for m in saw_machines if m['remark'] and not m['condition'])

add_kpi_card(slide6, 6.5, 1.2, 2.8, 1.4, 'Normal (ปกติ)', saw_ok, GREEN)
add_kpi_card(slide6, 9.6, 1.2, 2.8, 1.4, 'Dummy Tape OK', saw_dummy_ok, MED_BLUE)
add_kpi_card(slide6, 6.5, 2.9, 2.8, 1.4, 'Shutdown', saw_shutdown, RED)
add_kpi_card(slide6, 9.6, 2.9, 2.8, 1.4, 'Has Issues', saw_remark_issue, ORANGE)

# Issue list
add_text_box(slide6, 6.5, 4.6, 6.5, 0.35, 'Key Issues:', font_size=14, bold=True, color=DARK_BLUE)
issue_types = {
    'Vacuum fail (board discontinued)': sum(1 for m in saw_machines if 'vacuum' in m['remark'].lower()),
    'Spindle water leak': sum(1 for m in saw_machines if 'spindle' in m['remark'].lower()),
    'HDD fail (obsolete SAS)': sum(1 for m in saw_machines if 'HDD' in m['remark']),
    'CO2 Bubbler fail': sum(1 for m in saw_machines if 'CO2' in m['remark']),
    'Wait scrap': sum(1 for m in saw_machines if 'scrap' in m['remark'].lower()),
}
y_issue = 5.0
for issue, cnt in issue_types.items():
    if cnt > 0:
        add_text_box(slide6, 6.8, y_issue, 6, 0.30, f'\u2022  {issue}: {cnt} machines',
            font_size=12, color=RED if cnt > 2 else ORANGE)
        y_issue += 0.32

add_footer(slide6, 6)

# ========== SLIDE 7: Issues & Recommendations ==========
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide7, BG_LIGHT)
add_header_bar(slide7, 'Issues & Recommendations')

# Die Attach issues
add_shape_box(slide7, 0.5, 1.1, 6, 2.7, WHITE)
add_text_box(slide7, 0.7, 1.15, 5.6, 0.35, 'Die Attach Issues', font_size=16, bold=True, color=RED)
da_issues = [
    f'Boot Fail: {da_status_clean.get("Boot Fail",0)} machines - ELMO driver/parts unavailable',
    f'AD889 Turn-on: {upd_status_clean.get("Running",0)}/{len(update_machines)} running, {upd_status_clean.get("Boot Fail",0)} boot fail pending',
    f'Snap On Fail: 3 machines - PC mainboard damage (old model)',
    f'Wait PC MES: 1 machine - IS security update (~3 days)',
    f'ELMO driver shortage: {upd_status_clean.get("Boot Fail",0)} machines waiting for parts',
    f'On Stand By: {da_status_clean.get("On Stand By",0)} machines in inventory',
]
for i, item in enumerate(da_issues):
    add_text_box(slide7, 0.9, 1.55 + i*0.35, 5.4, 0.30, '\u2022  ' + item, font_size=11, color=BLACK)

# SAW issues
add_shape_box(slide7, 6.8, 1.1, 6, 2.7, WHITE)
add_text_box(slide7, 7.0, 1.15, 5.6, 0.35, 'SAW Machine Issues', font_size=16, bold=True, color=RED)
saw_issue_list = [
    f'Vacuum Fail: {issue_types.get("Vacuum fail (board discontinued)", 0)} machines - connector board discontinued',
    f'Spindle Water Leak: {issue_types.get("Spindle water leak", 0)} machines - spindle damage suspected',
    f'HDD Fail: {issue_types.get("HDD fail (obsolete SAS)", 0)} machines - obsolete SAS HDD, no replacement',
    f'CO2 Bubbler: {issue_types.get("CO2 Bubbler fail", 0)} machines shutdown - eFLOW unit failure',
    f'Wait Scrap: {issue_types.get("Wait scrap", 0)} machine - need to clear for LG#01',
]
for i, item in enumerate(saw_issue_list):
    add_text_box(slide7, 7.2, 1.55 + i*0.35, 5.4, 0.30, '\u2022  ' + item, font_size=11, color=BLACK)

# Recommendations
add_shape_box(slide7, 0.5, 4.1, 12.3, 3.0, WHITE)
add_text_box(slide7, 0.7, 4.15, 11.9, 0.35, 'Recommendations', font_size=16, bold=True, color=DARK_BLUE)
recs = [
    'Prioritize ELMO driver sourcing - Boot Fail is the #1 blocker for AD889 turn-on capacity',
    f'Source replacement for obsolete SAS HDD drives ({issue_types.get("HDD fail (obsolete SAS)", 0)} SAW machines affected)',
    'Evaluate repair/replacement options for vacuum connector boards (discontinued)',
    f'Assess spindle repair for {issue_types.get("Spindle water leak", 0)} SAW machines with water leak damage',
    'Resolve PC MES security update for DB#16 to enable additional AD889 capacity',
    f'Complete setup and monitoring for {upd_status_clean.get("Wait Setup/PC", 0) + upd_status_clean.get("On Going (Repair)", 0)} AD889 machines pending setup/repair',
    'Plan scrap disposal for DB#115 and SAW#015 to free up floor space',
    'Continue AD889 turn-on - target remaining machines once ELMO parts available',
]
for i, rec in enumerate(recs):
    add_text_box(slide7, 0.9, 4.55 + i*0.30, 11.5, 0.28, f'{i+1}. {rec}', font_size=11, color=BLACK)

add_footer(slide7, 7)

# Save
output_path = 'D:/claude/Project/Output/ASSY_CAP_Summary_Apr04_2026.pptx'
prs.save(output_path)
print(f'PPT saved to: {output_path}')
print(f'Total slides: {len(prs.slides)}')
print(f'Data: DA={da_total} machines, SAW={saw_total} machines, AD889 update={len(update_machines)} tracked')
print('Done!')
