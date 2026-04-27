import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"D:\claude\Project\raw\BOI 15 upgrade up Mar 20'26_finance confirm.xlsx"
OUT = r"D:\claude\Project\Output\BOI15_Assembly_Budget_Analysis.xlsx"

wb_src = openpyxl.load_workbook(SRC, data_only=True)

# ============ STYLES ============
DARK_BLUE = "1B3A5C"
MED_BLUE = "2E75B6"
GREEN = "2D8E4E"
RED = "C0392B"
ORANGE = "E67E22"
LIGHT_BLUE = "E8F0FE"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4E1"
LIGHT_ORANGE = "FFF2CC"
WHITE = "FFFFFF"

hdr_font = Font(name="Arial", bold=True, color=WHITE, size=11)
hdr_fill = PatternFill("solid", fgColor=DARK_BLUE)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
sub_hdr_font = Font(name="Arial", bold=True, color=WHITE, size=10)
sub_hdr_fill = PatternFill("solid", fgColor=MED_BLUE)
data_font = Font(name="Arial", size=10)
bold_font = Font(name="Arial", size=10, bold=True)
title_font = Font(name="Arial", bold=True, size=16, color=DARK_BLUE)
section_font = Font(name="Arial", bold=True, size=13, color=DARK_BLUE)
kpi_val_font = Font(name="Arial", bold=True, size=20, color=DARK_BLUE)
kpi_lbl_font = Font(name="Arial", size=9, color="666666")
thin_border = Border(
    left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0'))
alt_fill_1 = PatternFill("solid", fgColor=LIGHT_BLUE)
alt_fill_2 = PatternFill("solid", fgColor=WHITE)
green_fill = PatternFill("solid", fgColor=LIGHT_GREEN)
red_fill = PatternFill("solid", fgColor=LIGHT_RED)
orange_fill = PatternFill("solid", fgColor=LIGHT_ORANGE)

def style_header_row(ws, row, cols, font=hdr_font, fill=hdr_fill):
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = font
        cell.fill = fill
        cell.alignment = hdr_align
        cell.border = thin_border

def style_data_row(ws, row, cols, idx=0):
    fill = alt_fill_1 if idx % 2 == 0 else alt_fill_2
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = data_font
        cell.border = thin_border
        cell.fill = fill

def style_total_row(ws, row, cols, fill_color=MED_BLUE):
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.border = thin_border

# ============ PARSE SOURCE DATA ============
ws_upg = wb_src['Upgrade BOI10&11 to 15']
ws_pr = wb_src['PR']
ws_sum1m = wb_src['Sum 1M']
ws_summary = wb_src["Summary(P'TAN)"]

# --- Assembly Upgrade items ---
assy_upgrade = []
current_process = ''
for r in range(7, 79):
    b = ws_upg.cell(r, 2).value
    if not b:
        continue
    a = ws_upg.cell(r, 1).value
    if a and isinstance(a, str):
        current_process = a.strip()

    assy_upgrade.append({
        'process': current_process,
        'detail': b,
        'propose': ws_upg.cell(r, 3).value,
        'mc_no': ws_upg.cell(r, 6).value,
        'qty_set': ws_upg.cell(r, 7).value,
        'qty_mc': ws_upg.cell(r, 8).value,
        'cost_baht': ws_upg.cell(r, 9).value,
        'cost_usd': ws_upg.cell(r, 10).value,
        'total_baht': ws_upg.cell(r, 11).value,
        'total_usd': ws_upg.cell(r, 12).value or 0,
        'contact': ws_upg.cell(r, 13).value,
        'plant': ws_upg.cell(r, 14).value,
        'boi': ws_upg.cell(r, 15).value,
        'pr': ws_upg.cell(r, 17).value,
        'spent': ws_upg.cell(r, 18).value,
        'remain': ws_upg.cell(r, 19).value,
        'remark': ws_upg.cell(r, 20).value,
    })

# --- New Investment PRs ---
new_invest_prs = []
for r in range(3, ws_pr.max_row + 1):
    pr_num = ws_pr.cell(r, 4).value
    pr_usd = ws_pr.cell(r, 5).value
    if pr_num and pr_usd:
        new_invest_prs.append({'pr': str(pr_num), 'usd': pr_usd})

# --- Upgrade PRs ---
upgrade_prs = []
for r in range(3, ws_pr.max_row + 1):
    pr_num = ws_pr.cell(r, 1).value
    pr_usd = ws_pr.cell(r, 2).value
    if pr_num and pr_usd:
        upgrade_prs.append({'pr': str(pr_num), 'usd': pr_usd})

# --- Summary data ---
ni_depts = {}
for r in range(5, 14):
    dept = ws_summary.cell(r, 7).value
    actual = ws_summary.cell(r, 8).value or 0
    if dept:
        ni_depts[dept.strip()] = actual

# --- Sum 1M data ---
sum1m = {}
for r in [6, 7]:
    name = ws_sum1m.cell(r, 1).value
    sum1m[name] = {
        'upg_plan_k': ws_sum1m.cell(r, 2).value or 0,
        'upg_actual': ws_sum1m.cell(r, 3).value or 0,
        'upg_remain_k': ws_sum1m.cell(r, 5).value or 0,
        'ni_plan_k': ws_sum1m.cell(r, 6).value or 0,
        'ni_actual': ws_sum1m.cell(r, 7).value or 0,
        'ni_remain_k': ws_sum1m.cell(r, 9).value or 0,
        'grand_plan_k': ws_sum1m.cell(r, 10).value or 0,
    }

# ============ CREATE OUTPUT WORKBOOK ============
wb = Workbook()

# ============================
# SHEET 1: DASHBOARD
# ============================
ws1 = wb.active
ws1.title = "Assembly Dashboard"
ws1.sheet_properties.tabColor = DARK_BLUE

ws1.cell(1, 1, "BOI 15 - Assembly Budget Analysis").font = Font(name="Arial", bold=True, size=18, color=DARK_BLUE)
ws1.merge_cells("A1:J1")
ws1.cell(2, 1, "Upgrade + New Investment | MTAI & MMT | As of Mar 2026 (Finance Confirmed)").font = Font(name="Arial", size=11, color="666666")
ws1.merge_cells("A2:J2")

# --- KPI Section ---
r = 4
ws1.cell(r, 1, "ASSEMBLY BUDGET OVERVIEW ($K)").font = section_font
ws1.merge_cells(f"A{r}:J{r}")

# Table headers
r = 6
kpi_headers = ["Plant", "Type", "Budget Plan ($K)", "Actual Spent ($K)", "Remain ($K)", "% Used", "Status"]
for c, h in enumerate(kpi_headers, 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 7)

# MTAI data
r = 7
mtai = sum1m['Assembly MTHAI']
ws1.cell(r, 1, "MTAI")
ws1.cell(r, 2, "Upgrade")
ws1.cell(r, 3, mtai['upg_plan_k'])
ws1.cell(r, 4, mtai['upg_actual'] / 1000)
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_data_row(ws1, r, 7, 0)

r = 8
ws1.cell(r, 1, "MTAI")
ws1.cell(r, 2, "New Invest")
ws1.cell(r, 3, mtai['ni_plan_k'])
ws1.cell(r, 4, mtai['ni_actual'] / 1000)
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_data_row(ws1, r, 7, 1)

r = 9
ws1.cell(r, 1, "MTAI Total")
ws1.cell(r, 3).value = f'=C7+C8'
ws1.cell(r, 4).value = f'=D7+D8'
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_total_row(ws1, r, 7, MED_BLUE)

# MMT data
r = 10
mmt = sum1m['Assembly MMT']
ws1.cell(r, 1, "MMT")
ws1.cell(r, 2, "Upgrade")
ws1.cell(r, 3, mmt['upg_plan_k'])
ws1.cell(r, 4, mmt['upg_actual'] / 1000)
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_data_row(ws1, r, 7, 0)

r = 11
ws1.cell(r, 1, "MMT")
ws1.cell(r, 2, "New Invest")
ws1.cell(r, 3, mmt['ni_plan_k'])
ws1.cell(r, 4, mmt['ni_actual'] / 1000)
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_data_row(ws1, r, 7, 1)

r = 12
ws1.cell(r, 1, "MMT Total")
ws1.cell(r, 3).value = f'=C10+C11'
ws1.cell(r, 4).value = f'=D10+D11'
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_total_row(ws1, r, 7, MED_BLUE)

# Grand total
r = 14
ws1.cell(r, 1, "ASSEMBLY GRAND TOTAL")
ws1.cell(r, 3).value = '=C9+C12'
ws1.cell(r, 4).value = '=D9+D12'
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_total_row(ws1, r, 7, DARK_BLUE)

# Number formats for dashboard
for row in range(7, 15):
    for c in [3, 4, 5]:
        ws1.cell(row, c).number_format = '#,##0.00'
    ws1.cell(row, 6).number_format = '0.0%'

# --- Upgrade by Machine Type ---
r = 16
ws1.cell(r, 1, "ASSEMBLY UPGRADE BY MACHINE TYPE (Summary from P'TAN)").font = section_font
ws1.merge_cells(f"A{r}:G{r}")

r = 18
upg_headers = ["Machine Type", "Upgrade Cost (USD)", "Actual Spent (USD)", "Pending (USD)", "% Spent"]
for c, h in enumerate(upg_headers, 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 5)

# Read upgrade summary from Summary sheet
r = 19
upg_types = []
for sr in range(5, 23):
    mch = ws_summary.cell(sr, 2).value
    cost = ws_summary.cell(sr, 3).value
    actual = ws_summary.cell(sr, 4).value
    pending = ws_summary.cell(sr, 5).value
    if mch and cost:
        upg_types.append((mch, cost, actual or 0, pending or 0))

for idx, (mch, cost, actual, pending) in enumerate(upg_types):
    ws1.cell(r, 1, mch)
    ws1.cell(r, 2, cost)
    ws1.cell(r, 3, actual)
    ws1.cell(r, 4, pending)
    ws1.cell(r, 5).value = f'=IF(B{r}=0,"",C{r}/B{r})'
    ws1.cell(r, 5).number_format = '0.0%'
    for c in [2, 3, 4]:
        ws1.cell(r, c).number_format = '#,##0.00'
    style_data_row(ws1, r, 5, idx)
    r += 1

# Assy total row
assy_total_cost = ws_summary.cell(23, 3).value
assy_total_actual = ws_summary.cell(23, 4).value
assy_total_pending = ws_summary.cell(23, 5).value
ws1.cell(r, 1, "ASSY UPGRADE TOTAL")
ws1.cell(r, 2, assy_total_cost)
ws1.cell(r, 3, assy_total_actual)
ws1.cell(r, 4, assy_total_pending)
ws1.cell(r, 5).value = f'=IF(B{r}=0,"",C{r}/B{r})'
ws1.cell(r, 5).number_format = '0.0%'
for c in [2, 3, 4]:
    ws1.cell(r, c).number_format = '#,##0.00'
style_total_row(ws1, r, 5, DARK_BLUE)

# --- New Investment summary ---
r += 2
ws1.cell(r, 1, "ASSEMBLY NEW INVESTMENT (Summary)").font = section_font
ws1.merge_cells(f"A{r}:E{r}")

r += 2
ni_headers = ["Department", "Actual Spent (USD)", "Budget Plan ($K)", "Remain ($K)"]
for c, h in enumerate(ni_headers, 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 4)

r += 1
ws1.cell(r, 1, "ASSY MTHAI")
ws1.cell(r, 2, ni_depts.get('ASSY MTHAI', 0))
ws1.cell(r, 3, mtai['ni_plan_k'])
ws1.cell(r, 4).value = f'=C{r}-B{r}/1000'
for c in [2]: ws1.cell(r, c).number_format = '#,##0.00'
for c in [3, 4]: ws1.cell(r, c).number_format = '#,##0.00'
style_data_row(ws1, r, 4, 0)

r += 1
ws1.cell(r, 1, "ASSY MMT")
ws1.cell(r, 2, ni_depts.get('ASSY MMT', 0))
ws1.cell(r, 3, mmt['ni_plan_k'])
ws1.cell(r, 4).value = f'=C{r}-B{r}/1000'
for c in [2]: ws1.cell(r, c).number_format = '#,##0.00'
for c in [3, 4]: ws1.cell(r, c).number_format = '#,##0.00'
style_data_row(ws1, r, 4, 1)

r += 1
ws1.cell(r, 1, "ASSY NEW INVEST TOTAL")
ws1.cell(r, 2).value = f'=B{r-2}+B{r-1}'
ws1.cell(r, 3).value = f'=C{r-2}+C{r-1}'
ws1.cell(r, 4).value = f'=C{r}-B{r}/1000'
for c in [2]: ws1.cell(r, c).number_format = '#,##0.00'
for c in [3, 4]: ws1.cell(r, c).number_format = '#,##0.00'
style_total_row(ws1, r, 4, DARK_BLUE)

# Column widths
for c, w in {1:22, 2:20, 3:18, 4:18, 5:15, 6:12, 7:14}.items():
    ws1.column_dimensions[get_column_letter(c)].width = w
ws1.freeze_panes = "A3"

# ============================
# SHEET 2: ASSEMBLY UPGRADE DETAIL
# ============================
ws2 = wb.create_sheet("Assy Upgrade Detail")
ws2.sheet_properties.tabColor = MED_BLUE

ws2.cell(1, 1, "Assembly Upgrade Items - Full Detail").font = title_font
ws2.merge_cells("A1:P1")

headers2 = ["No.", "Plant", "Process", "Upgrade Details", "MC No.", "Qty/Set", "Qty MC",
            "Total (USD)", "Contact", "BOI", "PR#", "Spent (USD)", "Remain (USD)", "% Spent", "Status", "Remark"]
r = 3
for c, h in enumerate(headers2, 1):
    ws2.cell(r, c, h)
style_header_row(ws2, r, 16)

r = 4
# Sort: items with spending first, then by plant
assy_sorted = sorted(assy_upgrade, key=lambda x: (
    0 if (x['spent'] and float(x['spent'] or 0) > 0) else (1 if x['pr'] else 2),
    x['plant'] or '', x['process']))

for idx, item in enumerate(assy_sorted):
    spent = float(item['spent'] or 0)
    remain = item['remain']
    has_pr = item['pr'] is not None and str(item['pr']).strip() != ''

    if spent > 0 and (remain is None or float(remain) <= 0):
        status = "Completed"
    elif spent > 0:
        status = "In Progress"
    elif has_pr:
        status = "PR Issued"
    else:
        status = "Pending"

    ws2.cell(r, 1, idx + 1)
    ws2.cell(r, 2, item['plant'])
    ws2.cell(r, 3, item['process'])
    ws2.cell(r, 4, item['detail'])
    ws2.cell(r, 5, item['mc_no'])
    ws2.cell(r, 6, item['qty_set'])
    ws2.cell(r, 7, item['qty_mc'])
    ws2.cell(r, 8, item['total_usd'])
    ws2.cell(r, 9, item['contact'])
    ws2.cell(r, 10, str(item['boi']) if item['boi'] else '')
    ws2.cell(r, 11, item['pr'])
    ws2.cell(r, 12, item['spent'])
    ws2.cell(r, 13, item['remain'])
    ws2.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
    ws2.cell(r, 15, status)
    ws2.cell(r, 16, item['remark'])

    style_data_row(ws2, r, 16, idx)
    ws2.cell(r, 8).number_format = '#,##0.00'
    ws2.cell(r, 12).number_format = '#,##0.00'
    ws2.cell(r, 13).number_format = '#,##0.00'
    ws2.cell(r, 14).number_format = '0.0%'
    ws2.cell(r, 2).alignment = Alignment(horizontal="center")
    ws2.cell(r, 6).alignment = Alignment(horizontal="center")
    ws2.cell(r, 7).alignment = Alignment(horizontal="center")
    ws2.cell(r, 10).alignment = Alignment(horizontal="center")
    ws2.cell(r, 15).alignment = Alignment(horizontal="center")

    # Color status
    if status == "Completed":
        ws2.cell(r, 15).fill = PatternFill("solid", fgColor="C6EFCE")
        ws2.cell(r, 15).font = Font(name="Arial", size=10, color="006100")
    elif status == "In Progress":
        ws2.cell(r, 15).fill = PatternFill("solid", fgColor="FFEB9C")
        ws2.cell(r, 15).font = Font(name="Arial", size=10, color="9C6500")
    elif status == "PR Issued":
        ws2.cell(r, 15).fill = PatternFill("solid", fgColor="FCD5B4")
        ws2.cell(r, 15).font = Font(name="Arial", size=10, color="974706")

    r += 1

# Totals
last_data = r - 1
ws2.cell(r, 1, "TOTAL")
ws2.cell(r, 7).value = f'=SUM(G4:G{last_data})'
ws2.cell(r, 8).value = f'=SUM(H4:H{last_data})'
ws2.cell(r, 12).value = f'=SUM(L4:L{last_data})'
ws2.cell(r, 13).value = f'=SUM(M4:M{last_data})'
ws2.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
ws2.cell(r, 8).number_format = '#,##0.00'
ws2.cell(r, 12).number_format = '#,##0.00'
ws2.cell(r, 13).number_format = '#,##0.00'
ws2.cell(r, 14).number_format = '0.0%'
style_total_row(ws2, r, 16, DARK_BLUE)

# MTAI subtotal
r += 1
mtai_items = [i for i, item in enumerate(assy_sorted) if item['plant'] == 'MTAI']
ws2.cell(r, 1, "MTAI Subtotal")
ws2.cell(r, 8).value = sum(assy_sorted[i]['total_usd'] for i in mtai_items)
ws2.cell(r, 12).value = sum(float(assy_sorted[i]['spent'] or 0) for i in mtai_items)
ws2.cell(r, 13).value = sum(float(assy_sorted[i]['remain'] or 0) for i in mtai_items if assy_sorted[i]['remain'])
ws2.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
ws2.cell(r, 15, f"{len(mtai_items)} items")
for c in [8, 12, 13]: ws2.cell(r, c).number_format = '#,##0.00'
ws2.cell(r, 14).number_format = '0.0%'
style_total_row(ws2, r, 16, MED_BLUE)

# MMT subtotal
r += 1
mmt_items = [i for i, item in enumerate(assy_sorted) if item['plant'] == 'MMT']
ws2.cell(r, 1, "MMT Subtotal")
ws2.cell(r, 8).value = sum(assy_sorted[i]['total_usd'] for i in mmt_items)
ws2.cell(r, 12).value = sum(float(assy_sorted[i]['spent'] or 0) for i in mmt_items)
ws2.cell(r, 13).value = sum(float(assy_sorted[i]['remain'] or 0) for i in mmt_items if assy_sorted[i]['remain'])
ws2.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
ws2.cell(r, 15, f"{len(mmt_items)} items")
for c in [8, 12, 13]: ws2.cell(r, c).number_format = '#,##0.00'
ws2.cell(r, 14).number_format = '0.0%'
style_total_row(ws2, r, 16, MED_BLUE)

ws2.auto_filter.ref = f"A3:P{last_data}"
ws2.freeze_panes = "A4"

col_widths2 = {1:5, 2:7, 3:22, 4:55, 5:35, 6:8, 7:8, 8:13, 9:20, 10:8, 11:18, 12:13, 13:13, 14:9, 15:13, 16:30}
for c, w in col_widths2.items():
    ws2.column_dimensions[get_column_letter(c)].width = w

# ============================
# SHEET 3: NEW INVESTMENT DETAIL
# ============================
ws3 = wb.create_sheet("New Investment PRs")
ws3.sheet_properties.tabColor = GREEN

ws3.cell(1, 1, "New Investment - PR List (All Departments)").font = title_font
ws3.merge_cells("A1:F1")
ws3.cell(2, 1, "Assembly items highlighted | Source: PR sheet").font = Font(name="Arial", size=10, color="666666")

# Assembly NI summary
r = 4
ws3.cell(r, 1, "Assembly New Investment Summary").font = section_font
ws3.merge_cells(f"A{r}:D{r}")

r = 6
for c, h in enumerate(["Department", "Budget Plan ($K)", "Actual (USD)", "Remain ($K)"], 1):
    ws3.cell(r, c, h)
style_header_row(ws3, r, 4)

r = 7
ws3.cell(r, 1, "ASSY MTHAI")
ws3.cell(r, 2, mtai['ni_plan_k'])
ws3.cell(r, 3, ni_depts.get('ASSY MTHAI', 0))
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws3, r, 4, 0)
for c in [3]: ws3.cell(r, c).number_format = '#,##0.00'
for c in [2, 4]: ws3.cell(r, c).number_format = '#,##0.00'

r = 8
ws3.cell(r, 1, "ASSY MMT")
ws3.cell(r, 2, mmt['ni_plan_k'])
ws3.cell(r, 3, ni_depts.get('ASSY MMT', 0))
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws3, r, 4, 1)
for c in [3]: ws3.cell(r, c).number_format = '#,##0.00'
for c in [2, 4]: ws3.cell(r, c).number_format = '#,##0.00'

r = 9
ws3.cell(r, 1, "ASSY NI TOTAL")
ws3.cell(r, 2).value = f'=B7+B8'
ws3.cell(r, 3).value = f'=C7+C8'
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_total_row(ws3, r, 4, DARK_BLUE)
for c in [2, 3, 4]: ws3.cell(r, c).number_format = '#,##0.00'

# PR detail list
r = 11
ws3.cell(r, 1, "New Investment - PR Detail List").font = section_font
ws3.merge_cells(f"A{r}:D{r}")

r = 13
for c, h in enumerate(["No.", "PR#", "Amount (USD)", "Note"], 1):
    ws3.cell(r, c, h)
style_header_row(ws3, r, 4)

r = 14
# known ASSY MTHAI PR
assy_mthai_pr = '100104009'  # exact match $196,793.07
for idx, p in enumerate(new_invest_prs):
    ws3.cell(r, 1, idx + 1)
    ws3.cell(r, 2, p['pr'])
    ws3.cell(r, 3, p['usd'])
    ws3.cell(r, 3).number_format = '#,##0.00'
    if p['pr'] == assy_mthai_pr:
        ws3.cell(r, 4, "ASSY MTHAI (confirmed)")
        for c in range(1, 5):
            ws3.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    style_data_row(ws3, r, 4, idx)
    # Re-apply green highlight for known ASSY PRs
    if p['pr'] == assy_mthai_pr:
        for c in range(1, 5):
            ws3.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            ws3.cell(r, c).font = Font(name="Arial", size=10, bold=True, color="006100")
    r += 1

# Total
ws3.cell(r, 1, "TOTAL")
ws3.cell(r, 3).value = f'=SUM(C14:C{r-1})'
ws3.cell(r, 3).number_format = '#,##0.00'
style_total_row(ws3, r, 4, DARK_BLUE)

ws3.column_dimensions['A'].width = 6
ws3.column_dimensions['B'].width = 18
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 30
ws3.freeze_panes = "A3"

# ============================
# SHEET 4: ALL UPGRADE PRs
# ============================
ws4 = wb.create_sheet("Upgrade PRs")
ws4.sheet_properties.tabColor = ORANGE

ws4.cell(1, 1, "Upgrade - PR List (All Departments)").font = title_font
ws4.merge_cells("A1:E1")

# Map PR to item detail from upgrade sheet for ASSY items
pr_to_item = {}
for item in assy_upgrade:
    if item['pr']:
        # PR field may have multiple PRs comma-separated
        for pr_raw in str(item['pr']).split(','):
            pr_clean = pr_raw.strip()
            if pr_clean:
                pr_to_item[pr_clean] = f"{item['process']} - {str(item['detail'])[:40]}"

r = 3
for c, h in enumerate(["No.", "PR#", "Amount (USD)", "Linked Item (if Assembly)", "Type"], 1):
    ws4.cell(r, c, h)
style_header_row(ws4, r, 5)

r = 4
for idx, p in enumerate(upgrade_prs):
    ws4.cell(r, 1, idx + 1)
    ws4.cell(r, 2, p['pr'])
    ws4.cell(r, 3, p['usd'])
    ws4.cell(r, 3).number_format = '#,##0.00'
    linked = pr_to_item.get(p['pr'], '')
    ws4.cell(r, 4, linked)
    ws4.cell(r, 5, "ASSY" if linked else "TEST/Other")
    style_data_row(ws4, r, 5, idx)
    if linked:
        for c in range(1, 6):
            ws4.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    r += 1

ws4.cell(r, 1, "TOTAL")
ws4.cell(r, 3).value = f'=SUM(C4:C{r-1})'
ws4.cell(r, 3).number_format = '#,##0.00'
style_total_row(ws4, r, 5, DARK_BLUE)

ws4.auto_filter.ref = f"A3:E{r-1}"
ws4.freeze_panes = "A4"
ws4.column_dimensions['A'].width = 6
ws4.column_dimensions['B'].width = 18
ws4.column_dimensions['C'].width = 18
ws4.column_dimensions['D'].width = 50
ws4.column_dimensions['E'].width = 14

# ============ SAVE ============
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Sheets: {wb.sheetnames}")
print(f"Assembly upgrade items: {len(assy_upgrade)}")
print(f"  MTAI: {len([i for i in assy_upgrade if i['plant']=='MTAI'])}")
print(f"  MMT:  {len([i for i in assy_upgrade if i['plant']=='MMT'])}")
print(f"New Investment PRs: {len(new_invest_prs)}")
print(f"Upgrade PRs: {len(upgrade_prs)}")
