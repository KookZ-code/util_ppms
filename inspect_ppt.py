# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400

def emu_to_inches(emu):
    if emu is None:
        return None
    return round(emu / EMU_PER_INCH, 4)

def rgb_to_hex(rgb):
    if rgb is None:
        return None
    return f"#{rgb.r:02X}{rgb.g:02X}{rgb.b:02X}"

def inspect_text_frame(tf, indent="    "):
    for i, para in enumerate(tf.paragraphs):
        full_text = para.text
        alignment = para.alignment
        para_font_size = None
        para_font_bold = None
        para_font_color = None
        para_font_name = None

        if para.runs:
            for run in para.runs:
                f = run.font
                if f.size:
                    para_font_size = f.size.pt
                if f.bold is not None:
                    para_font_bold = f.bold
                try:
                    if f.color and f.color.type is not None:
                        para_font_color = rgb_to_hex(f.color.rgb)
                except:
                    pass
                if f.name:
                    para_font_name = f.name

        if full_text.strip() or i == 0:
            print(f"{indent}Para[{i}]: '{full_text}'")
            print(f"{indent}       align={alignment}, font_size={para_font_size}pt, bold={para_font_bold}, color={para_font_color}, font={para_font_name}")

def inspect_shape(shape, indent="  "):
    shape_type_name = str(shape.shape_type)
    print(f"{indent}Shape: '{shape.name}' | Type: {shape_type_name}")
    print(f"{indent}  Pos: L={emu_to_inches(shape.left)}\", T={emu_to_inches(shape.top)}\"")
    print(f"{indent}  Size: W={emu_to_inches(shape.width)}\", H={emu_to_inches(shape.height)}\"")

    # Fill info
    try:
        fill = shape.fill
        fill_type = fill.type
        fill_color = None
        if fill_type is not None:
            try:
                fill_color = rgb_to_hex(fill.fore_color.rgb)
            except:
                pass
        print(f"{indent}  Fill type={fill_type}, color={fill_color}")
    except:
        print(f"{indent}  Fill: (error reading)")

    # Line/border info
    try:
        line = shape.line
        line_color = None
        try:
            line_color = rgb_to_hex(line.color.rgb)
        except:
            pass
        print(f"{indent}  Line: width={line.width}, color={line_color}")
    except:
        print(f"{indent}  Line: (error reading)")

    # Text frame
    if shape.has_text_frame:
        tf = shape.text_frame
        print(f"{indent}  TextFrame: word_wrap={tf.word_wrap}, auto_size={tf.auto_size}")
        try:
            anchor = tf.vertical_anchor
            print(f"{indent}  TextFrame anchor={anchor}")
        except:
            pass
        inspect_text_frame(tf, indent + "    ")

    # Table
    if shape.has_table:
        table = shape.table
        print(f"{indent}  Table: {table.rows.__len__()} rows x {len(table.columns)} cols")
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text_frame.text
                print(f"{indent}    Cell[{r_idx},{c_idx}]: '{cell_text}'")

    # Chart
    if shape.has_chart:
        chart = shape.chart
        print(f"{indent}  Chart type: {chart.chart_type}")
        for series in chart.series:
            print(f"{indent}    Series: '{series.name}'")

    # Picture
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        print(f"{indent}  [Picture/Image]")

def main():
    pptx_path = r"D:\claude\Project\Output\DA_ASSY_CAP_Exec_Summary_Apr22.pptx"
    prs = Presentation(pptx_path)

    print(f"Presentation: {pptx_path}")
    print(f"Slide size: {emu_to_inches(prs.slide_width)}\" x {emu_to_inches(prs.slide_height)}\"")
    print(f"Total slides: {len(prs.slides)}")
    print("=" * 80)

    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        print(f"\n{'='*80}")
        print(f"SLIDE {slide_num}")
        print(f"{'='*80}")

        # Background
        try:
            bg = slide.background
            fill = bg.fill
            print(f"  Background fill type: {fill.type}")
            try:
                print(f"  Background color: {rgb_to_hex(fill.fore_color.rgb)}")
            except:
                print(f"  Background color: (none or theme)")
        except:
            print(f"  Background: (error)")

        print(f"\n  Shapes ({len(slide.shapes)}):")
        for shape in slide.shapes:
            inspect_shape(shape, "  ")
            print()

if __name__ == '__main__':
    main()
