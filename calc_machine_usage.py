import sys, io, re, math
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb_src = openpyxl.load_workbook('raw/STD UPH vs DRR.xlsx', data_only=True)

# ── Sheet 1: UPH ──
ws1 = wb_src['UPH FOL _ EOL Req ']
processes = []
for c in range(2, ws1.max_column + 1):
    val = ws1.cell(1, c).value
    processes.append(val.replace('UPH_', '') if val else f'Process{c}')

uph_data = {}
for r in range(2, ws1.max_row + 1):
    pkg = ws1.cell(r, 1).value
    if not pkg:
        continue
    pkg = pkg.strip()
    uphs = []
    for c in range(2, ws1.max_column + 1):
        v = ws1.cell(r, c).value
        uphs.append(v if v else 0)
    uph_data[pkg] = uphs

# ── Sheet 2: DDR ──
ws2 = wb_src['DDR ww02']
ddr_data = []
for r in range(3, ws2.max_row + 1):
    name = ws2.cell(r, 1).value
    val = ws2.cell(r, 2).value
    if name and val is not None:
        ddr_data.append((name.strip(), val))

# ── Matching function ──
def match_ddr_to_uph(name):
    clean = re.sub(r'\([^)]*\)[WP]*$', '', name).strip()
    clean = re.sub(r'\([^)]*\)', '', clean).strip()
    clean = re.sub(r'_\w+$', '', clean).strip()
    m = re.match(r'^(\d+)', clean)
    if not m:
        return None
    leads = m.group(1)
    rest = clean[m.end():].strip()
    rest = re.sub(r'^L\s*', '', rest).strip()
    variant = ''
    for suf in [' HD', ' UD', ' IDF']:
        if rest.endswith(suf):
            variant = suf.strip()
            rest = rest[:-len(suf)].strip()
            break
    size_match = re.search(r'(\d+(?:\.\d+)?)\s*[xX]\s*(\d+(?:\.\d+)?)', rest)
    if size_match:
        size = f'{size_match.group(1)}X{size_match.group(2)}'.upper()
        pkgtype = rest[:size_match.start()].strip()
    else:
        size = ''
        pkgtype = rest.strip()
    if variant == 'HD':
        cand = f'{leads}{pkgtype}-HD'
    elif variant == 'UD':
        cand = f'{leads}{pkgtype}-HD'
    else:
        cand = f'{leads}{pkgtype} {size}'.strip() if size else f'{leads}{pkgtype}'
    if cand in uph_data:
        return cand
    if 'TQFP' in pkgtype and not size:
        for p in uph_data:
            if p.startswith(f'{leads}TQFP') and '-HD' not in p:
                return p
    if 'EIAJ' in pkgtype:
        cand2 = f'{leads}SOIJ'
        if cand2 in uph_data:
            return cand2
    return None

# ── Build result rows (only DDR > 0 with match) ──
results = []
for ddr_name, ddr_k in ddr_data:
    if ddr_k == 0:
        continue
    matched = match_ddr_to_uph(ddr_name)
    if not matched:
        continue
    results.append({
        'ddr_name': ddr_name,
        'uph_pkg': matched,
        'ddr_k': ddr_k,
        'uphs': uph_data[matched],
    })

# ── Styles ──
hdr_fill = PatternFill('solid', fgColor='1B3A5C')
hdr_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
sub_fill = PatternFill('solid', fgColor='D6E4F0')
sub_font = Font(name='Calibri', size=10, bold=True, color='1B3A5C')
num_font = Font(name='Calibri', size=10)
input_fill = PatternFill('solid', fgColor='FFF3CD')       # yellow for editable
input_hdr_fill = PatternFill('solid', fgColor='E67E22')   # orange header
input_hdr_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
green_fill = PatternFill('solid', fgColor='E8F5E9')
yellow_fill = PatternFill('solid', fgColor='FFF8E1')
red_fill = PatternFill('solid', fgColor='FFEBEE')
thin_border = Border(
    left=Side(style='thin', color='B0B0B0'),
    right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='B0B0B0'),
    bottom=Side(style='thin', color='B0B0B0'),
)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_al = Alignment(horizontal='left', vertical='center', wrap_text=True)

def style_cell(cell, font=num_font, fill=None, align=center):
    cell.font = font
    cell.border = thin_border
    cell.alignment = align
    if fill:
        cell.fill = fill

# ── Create output workbook ──
wb_out = openpyxl.Workbook()
N = len(results)          # number of data rows
first_data = 2            # first data row
last_data = N + 1         # last data row
total_row = N + 2         # total row

# ================================================================
# SHEET: UPH Ref  (hidden reference for formulas)
# ================================================================
ws_ref = wb_out.active
ws_ref.title = 'UPH Ref'

# Header
ws_ref.cell(1, 1, 'PKG')
for j, p in enumerate(processes):
    ws_ref.cell(1, 2 + j, p)

# Data — same row order as results
for i, res in enumerate(results):
    r = i + 2
    ws_ref.cell(r, 1, res['uph_pkg'])
    for j, u in enumerate(res['uphs']):
        ws_ref.cell(r, 2 + j, u)

ws_ref.sheet_state = 'hidden'

# ================================================================
# SHEET: Machine Hours Detail
# ================================================================
# Columns:
#   A(1): DDR Product
#   B(2): Matched UPH PKG
#   C(3): DDR WW02 (K units)
#   D(4): Future DDR (K units)   ← EDITABLE INPUT
#   E(5): Active DDR (units)     ← formula
#   F(6)..Q(17): Hrs per process ← formulas

ws_det = wb_out.create_sheet('Machine Hours Detail')
wb_out.move_sheet('Machine Hours Detail', offset=-1)  # move before UPH Ref

det_headers = ['DDR Product', 'Matched UPH PKG', 'DDR WW02\n(K units)',
               'Future DDR\n(K units)', 'Active DDR\n(units)']
for p in processes:
    det_headers.append(f'Hrs\n{p}')
max_col_det = len(det_headers)

# Write headers
for c, h in enumerate(det_headers, 1):
    cell = ws_det.cell(1, c, h)
    if c == 4:
        style_cell(cell, font=input_hdr_font, fill=input_hdr_fill)
    else:
        style_cell(cell, font=hdr_font, fill=hdr_fill)

# Write data rows with FORMULAS
for i, res in enumerate(results):
    r = i + 2
    # A: product name
    cell_a = ws_det.cell(r, 1, res['ddr_name'])
    style_cell(cell_a, align=left_al)
    # B: UPH package
    cell_b = ws_det.cell(r, 2, res['uph_pkg'])
    style_cell(cell_b, align=left_al)
    # C: current DDR (K)
    cell_c = ws_det.cell(r, 3, res['ddr_k'])
    style_cell(cell_c)
    cell_c.number_format = '#,##0'
    # D: Future DDR (K) — EMPTY, editable
    cell_d = ws_det.cell(r, 4)
    style_cell(cell_d, fill=input_fill)
    cell_d.number_format = '#,##0'
    # E: Active DDR (units) = IF(D="", C, D) * 1000
    cell_e = ws_det.cell(r, 5)
    cell_e.value = f'=IF(D{r}="",C{r},D{r})*1000'
    style_cell(cell_e)
    cell_e.number_format = '#,##0'
    # F..Q: Machine hours per process
    # Formula: =IF('UPH Ref'!{col}{r}=0, "", $E{r}/'UPH Ref'!{col}{r})
    for j in range(len(processes)):
        uph_col = get_column_letter(2 + j)  # UPH Ref columns B..M
        cell = ws_det.cell(r, 6 + j)
        cell.value = f"=IF('UPH Ref'!{uph_col}{r}=0,\"\",$E{r}/'UPH Ref'!{uph_col}{r})"
        style_cell(cell)
        cell.number_format = '#,##0.00'

# TOTAL row
ws_det.cell(total_row, 1, 'TOTAL')
style_cell(ws_det.cell(total_row, 1), font=sub_font, fill=sub_fill, align=left_al)
for c in [2]:
    style_cell(ws_det.cell(total_row, c), font=sub_font, fill=sub_fill)
# C total: SUM of current DDR
ws_det.cell(total_row, 3).value = f'=SUM(C{first_data}:C{last_data})'
ws_det.cell(total_row, 3).number_format = '#,##0'
style_cell(ws_det.cell(total_row, 3), font=sub_font, fill=sub_fill)
# D total: SUM of future DDR (if any filled)
ws_det.cell(total_row, 4).value = f'=IF(COUNTA(D{first_data}:D{last_data})=0,"",SUM(D{first_data}:D{last_data}))'
ws_det.cell(total_row, 4).number_format = '#,##0'
style_cell(ws_det.cell(total_row, 4), font=sub_font, fill=input_fill)
# E total: SUM of active DDR
ws_det.cell(total_row, 5).value = f'=SUM(E{first_data}:E{last_data})'
ws_det.cell(total_row, 5).number_format = '#,##0'
style_cell(ws_det.cell(total_row, 5), font=sub_font, fill=sub_fill)
# Process totals: SUM
for j in range(len(processes)):
    col_letter = get_column_letter(6 + j)
    cell = ws_det.cell(total_row, 6 + j)
    cell.value = f'=SUM({col_letter}{first_data}:{col_letter}{last_data})'
    cell.number_format = '#,##0.00'
    style_cell(cell, font=sub_font, fill=sub_fill)

# Column widths
ws_det.column_dimensions['A'].width = 35
ws_det.column_dimensions['B'].width = 22
ws_det.column_dimensions['C'].width = 15
ws_det.column_dimensions['D'].width = 17
ws_det.column_dimensions['E'].width = 16
for c in range(6, max_col_det + 1):
    ws_det.column_dimensions[get_column_letter(c)].width = 16
ws_det.freeze_panes = 'F2'
ws_det.sheet_properties.tabColor = '2E75B6'

# ================================================================
# SHEET: Summary by Process  (all formulas)
# ================================================================
ws_sum = wb_out.create_sheet('Summary by Process')

DET_SHEET = "'Machine Hours Detail'"
sum_headers = ['Process', 'Total Machine\nHrs/Day', 'Machines Needed\n(21hr)',
               'Machines Needed\n(20hr)', 'Machines Needed\n(16hr)']
for c, h in enumerate(sum_headers, 1):
    cell = ws_sum.cell(1, c, h)
    style_cell(cell, font=hdr_font, fill=hdr_fill)

for j, proc in enumerate(processes):
    r = j + 2
    proc_col = get_column_letter(6 + j)  # F..Q in detail sheet
    ws_sum.cell(r, 1, proc)
    style_cell(ws_sum.cell(r, 1), align=left_al)
    # B: total hrs = reference detail total row
    ws_sum.cell(r, 2).value = f"={DET_SHEET}!{proc_col}{total_row}"
    ws_sum.cell(r, 2).number_format = '#,##0.00'
    style_cell(ws_sum.cell(r, 2))
    # C: machines at 21hr
    ws_sum.cell(r, 3).value = f'=IF(B{r}=0,0,ROUNDUP(B{r}/21,0))'
    style_cell(ws_sum.cell(r, 3))
    # D: machines at 20hr
    ws_sum.cell(r, 4).value = f'=IF(B{r}=0,0,ROUNDUP(B{r}/20,0))'
    style_cell(ws_sum.cell(r, 4))
    # E: machines at 16hr
    ws_sum.cell(r, 5).value = f'=IF(B{r}=0,0,ROUNDUP(B{r}/16,0))'
    style_cell(ws_sum.cell(r, 5))

# Total row for summary
sum_total = len(processes) + 2
ws_sum.cell(sum_total, 1, 'TOTAL')
style_cell(ws_sum.cell(sum_total, 1), font=sub_font, fill=sub_fill, align=left_al)
for c in range(2, 6):
    col_l = get_column_letter(c)
    ws_sum.cell(sum_total, c).value = f'=SUM({col_l}2:{col_l}{sum_total-1})'
    ws_sum.cell(sum_total, c).number_format = '#,##0.00' if c == 2 else '#,##0'
    style_cell(ws_sum.cell(sum_total, c), font=sub_font, fill=sub_fill)

ws_sum.column_dimensions['A'].width = 28
for c in range(2, 6):
    ws_sum.column_dimensions[get_column_letter(c)].width = 20
ws_sum.freeze_panes = 'B2'
ws_sum.sheet_properties.tabColor = '2D8E4E'

# ================================================================
# SHEET: By UPH Package  (SUMIF formulas)
# ================================================================
ws_pkg = wb_out.create_sheet('By UPH Package')

# Collect unique UPH packages (sorted)
pkg_agg = {}
for res in results:
    pkg = res['uph_pkg']
    if pkg not in pkg_agg:
        pkg_agg[pkg] = []
    pkg_agg[pkg].append(res['ddr_name'])
sorted_pkgs = sorted(pkg_agg.keys())

pkg_headers = ['UPH Package', 'DDR Products', 'Total Active\nDDR (K)']
for p in processes:
    pkg_headers.append(f'Hrs\n{p}')
max_col_pkg = len(pkg_headers)

for c, h in enumerate(pkg_headers, 1):
    cell = ws_pkg.cell(1, c, h)
    style_cell(cell, font=hdr_font, fill=hdr_fill)

for i, pkg in enumerate(sorted_pkgs):
    r = i + 2
    ws_pkg.cell(r, 1, pkg)
    style_cell(ws_pkg.cell(r, 1), align=left_al)
    ws_pkg.cell(r, 2, ', '.join(pkg_agg[pkg]))
    style_cell(ws_pkg.cell(r, 2), align=left_al)
    # C: Total active DDR (K) = SUMIF on detail col B matching pkg, sum col E / 1000
    ws_pkg.cell(r, 3).value = (
        f"=SUMIF({DET_SHEET}!$B${first_data}:$B${last_data},$A{r},{DET_SHEET}!$E${first_data}:$E${last_data})/1000"
    )
    ws_pkg.cell(r, 3).number_format = '#,##0'
    style_cell(ws_pkg.cell(r, 3))
    # D onwards: Hrs per process = SUMIF
    for j in range(len(processes)):
        proc_col = get_column_letter(6 + j)
        cell = ws_pkg.cell(r, 4 + j)
        cell.value = (
            f"=SUMIF({DET_SHEET}!$B${first_data}:$B${last_data},$A{r},"
            f"{DET_SHEET}!${proc_col}${first_data}:${proc_col}${last_data})"
        )
        cell.number_format = '#,##0.00'
        style_cell(cell)

# Totals
pkg_total = len(sorted_pkgs) + 2
ws_pkg.cell(pkg_total, 1, 'TOTAL')
style_cell(ws_pkg.cell(pkg_total, 1), font=sub_font, fill=sub_fill, align=left_al)
style_cell(ws_pkg.cell(pkg_total, 2), font=sub_font, fill=sub_fill)
for c in range(3, max_col_pkg + 1):
    col_l = get_column_letter(c)
    ws_pkg.cell(pkg_total, c).value = f'=SUM({col_l}2:{col_l}{pkg_total-1})'
    ws_pkg.cell(pkg_total, c).number_format = '#,##0.00' if c > 3 else '#,##0'
    style_cell(ws_pkg.cell(pkg_total, c), font=sub_font, fill=sub_fill)

ws_pkg.column_dimensions['A'].width = 22
ws_pkg.column_dimensions['B'].width = 45
ws_pkg.column_dimensions['C'].width = 16
for c in range(4, max_col_pkg + 1):
    ws_pkg.column_dimensions[get_column_letter(c)].width = 16
ws_pkg.freeze_panes = 'D2'

# ================================================================
# SHEET: Reverse Calc — Input machines → max supportable loading
# ================================================================
ws_rev = wb_out.create_sheet('Reverse Calc')

SUM_SHEET = "'Summary by Process'"
title_font = Font(name='Calibri', size=14, bold=True, color='1B3A5C')
label_font = Font(name='Calibri', size=11, bold=True, color='1B3A5C')
result_fill = PatternFill('solid', fgColor='D5F5E3')
result_font = Font(name='Calibri', size=12, bold=True, color='1B3A5C')
bottleneck_fill = PatternFill('solid', fgColor='FADBD8')
bottleneck_font = Font(name='Calibri', size=12, bold=True, color='C0392B')

# Row 1: Title
ws_rev.merge_cells('A1:H1')
c1 = ws_rev.cell(1, 1, 'Reverse Calculation — Input # Machines → Max Supportable Loading')
c1.font = title_font
c1.alignment = Alignment(horizontal='left', vertical='center')

# Row 2: Working hours input
ws_rev.cell(2, 1, 'Working Hours / Day :')
ws_rev.cell(2, 1).font = label_font
c_hrs = ws_rev.cell(2, 2, 21)
style_cell(c_hrs, fill=input_fill)
c_hrs.number_format = '0'
ws_rev.cell(2, 3, '(editable)')
ws_rev.cell(2, 3).font = Font(name='Calibri', size=9, italic=True, color='888888')

# Row 3: blank
# Row 4: Headers
P_FIRST = 5                      # first process data row
P_LAST = P_FIRST + len(processes) - 1  # last process data row
rev_headers = [
    'Process',
    '# Machines\n(Input)',
    'Available\nHrs/Day',
    'Current Hrs\nNeeded/Day',
    'Utilization %',
    'Remaining\nHrs',
    'Scale Factor',
    'Max Supportable\nDDR (K units)',
]
for c, h in enumerate(rev_headers, 1):
    cell = ws_rev.cell(4, c, h)
    if c == 2:
        style_cell(cell, font=input_hdr_font, fill=input_hdr_fill)
    else:
        style_cell(cell, font=hdr_font, fill=hdr_fill)

# Rows 5..16: Process data
for j, proc in enumerate(processes):
    r = P_FIRST + j
    sum_r = 2 + j  # corresponding row in Summary sheet

    # A: Process name
    ws_rev.cell(r, 1, proc)
    style_cell(ws_rev.cell(r, 1), align=left_al)

    # B: # Machines — EDITABLE INPUT (yellow)
    style_cell(ws_rev.cell(r, 2), fill=input_fill)
    ws_rev.cell(r, 2).number_format = '#,##0'

    # C: Available Hrs = machines × working_hours
    ws_rev.cell(r, 3).value = f'=IF(B{r}="","",B{r}*$B$2)'
    ws_rev.cell(r, 3).number_format = '#,##0.0'
    style_cell(ws_rev.cell(r, 3))

    # D: Current Hrs Needed (from Summary sheet)
    ws_rev.cell(r, 4).value = f"={SUM_SHEET}!B{sum_r}"
    ws_rev.cell(r, 4).number_format = '#,##0.00'
    style_cell(ws_rev.cell(r, 4))

    # E: Utilization % = D / C
    ws_rev.cell(r, 5).value = f'=IF(OR(C{r}="",C{r}=0),"",D{r}/C{r})'
    ws_rev.cell(r, 5).number_format = '0.0%'
    style_cell(ws_rev.cell(r, 5))

    # F: Remaining Hrs = C - D
    ws_rev.cell(r, 6).value = f'=IF(C{r}="","",C{r}-D{r})'
    ws_rev.cell(r, 6).number_format = '#,##0.00'
    style_cell(ws_rev.cell(r, 6))

    # G: Scale Factor = C / D (how much current load can be scaled)
    ws_rev.cell(r, 7).value = f'=IF(OR(D{r}=0,C{r}=""),"",C{r}/D{r})'
    ws_rev.cell(r, 7).number_format = '0.00x'  # custom format
    style_cell(ws_rev.cell(r, 7))

    # H: Max DDR this process supports = current total DDR × scale factor
    #    Current Total DDR is at row RS+1 = P_LAST+3
    ws_rev.cell(r, 8).value = f'=IF(G{r}="","",G{r}*$B${P_LAST + 3})'
    ws_rev.cell(r, 8).number_format = '#,##0'
    style_cell(ws_rev.cell(r, 8))

# ── Result Summary section ──
RS = P_LAST + 2  # blank row after processes, then results

ws_rev.merge_cells(f'A{RS}:H{RS}')
ws_rev.cell(RS, 1, 'RESULT SUMMARY')
ws_rev.cell(RS, 1).font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
ws_rev.cell(RS, 1).fill = PatternFill('solid', fgColor='1B3A5C')
ws_rev.cell(RS, 1).alignment = Alignment(horizontal='center', vertical='center')
for cc in range(2, 9):
    ws_rev.cell(RS, cc).fill = PatternFill('solid', fgColor='1B3A5C')

# Row RS+1: Current Total DDR
r1 = RS + 1
ws_rev.cell(r1, 1, 'Current Active Total DDR (K units)')
ws_rev.cell(r1, 1).font = label_font
ws_rev.cell(r1, 2).value = f"={DET_SHEET}!E{total_row}/1000"
ws_rev.cell(r1, 2).number_format = '#,##0'
style_cell(ws_rev.cell(r1, 2), font=result_font, fill=result_fill)

# Row RS+2: Bottleneck Process (lowest scale factor)
r2 = RS + 2
ws_rev.cell(r2, 1, 'Bottleneck Process')
ws_rev.cell(r2, 1).font = label_font
ws_rev.cell(r2, 2).value = (
    f'=IFERROR(INDEX(A{P_FIRST}:A{P_LAST},'
    f'MATCH(MIN(G{P_FIRST}:G{P_LAST}),G{P_FIRST}:G{P_LAST},0)),"")'
)
style_cell(ws_rev.cell(r2, 2), font=bottleneck_font, fill=bottleneck_fill, align=left_al)

# Row RS+3: Min Scale Factor
r3 = RS + 3
ws_rev.cell(r3, 1, 'Min Scale Factor (bottleneck)')
ws_rev.cell(r3, 1).font = label_font
ws_rev.cell(r3, 2).value = f'=IFERROR(MIN(G{P_FIRST}:G{P_LAST}),"")'
ws_rev.cell(r3, 2).number_format = '0.00x'
style_cell(ws_rev.cell(r3, 2), font=result_font, fill=result_fill)

# Row RS+4: Max Supportable Total DDR
r4 = RS + 4
ws_rev.cell(r4, 1, 'Max Supportable Total DDR (K units)')
ws_rev.cell(r4, 1).font = Font(name='Calibri', size=12, bold=True, color='2D8E4E')
ws_rev.cell(r4, 2).value = f'=IFERROR(B{r1}*B{r3},"")'
ws_rev.cell(r4, 2).number_format = '#,##0'
style_cell(ws_rev.cell(r4, 2),
           font=Font(name='Calibri', size=14, bold=True, color='2D8E4E'),
           fill=PatternFill('solid', fgColor='D5F5E3'))

# Row RS+6: Note
r_note = RS + 6
ws_rev.cell(r_note, 1, 'Note: Assumes same product mix ratio as current Active DDR. '
            'WB-Copper UPH is per-head, not per-machine.')
ws_rev.cell(r_note, 1).font = Font(name='Calibri', size=9, italic=True, color='888888')
ws_rev.merge_cells(f'A{r_note}:H{r_note}')

# Column widths
ws_rev.column_dimensions['A'].width = 36
ws_rev.column_dimensions['B'].width = 16
ws_rev.column_dimensions['C'].width = 16
ws_rev.column_dimensions['D'].width = 18
ws_rev.column_dimensions['E'].width = 15
ws_rev.column_dimensions['F'].width = 15
ws_rev.column_dimensions['G'].width = 15
ws_rev.column_dimensions['H'].width = 22
ws_rev.freeze_panes = 'A5'
ws_rev.sheet_properties.tabColor = 'E67E22'

# ================================================================
# SHEET: Notes
# ================================================================
ws_note = wb_out.create_sheet('Notes')
notes = [
    ['Machine Usage Calculation — Notes'],
    [''],
    ['Formula: Machine Hours = Active DDR (units) / Standard UPH'],
    ['Active DDR (units) = IF(Future DDR is filled, Future DDR, Current DDR WW02) × 1000'],
    ['Machines Needed = Machine Hours / Available Hours per Day (rounded up)'],
    [''],
    ['HOW TO USE:'],
    ['- To simulate future load changes, enter new DDR values (K units) in the yellow "Future DDR" column'],
    ['- All machine hours, summaries, and package aggregations update automatically via formulas'],
    ['- Leave a cell blank in "Future DDR" to keep using the current WW02 value'],
    ['- The "UPH Ref" sheet (hidden) contains standard UPH data used by formulas'],
    [''],
    ['REVERSE CALC SHEET:'],
    ['- Input the number of machines per process in column B (yellow cells)'],
    ['- The sheet calculates available hours, utilization, and scale factor per process'],
    ['- The bottleneck process (lowest scale factor) determines max total loading'],
    ['- "Max Supportable DDR" = Current Total DDR × Min Scale Factor'],
    ['- Assumes the same product mix ratio as current Active DDR'],
    [''],
    ['DDR Source: Sheet "DDR ww02" from STD UPH vs DRR.xlsx'],
    ['UPH Source: Sheet "UPH FOL _ EOL Req" from STD UPH vs DRR.xlsx'],
    [''],
    ['IMPORTANT NOTES:'],
    ['- WB-Copper (Wire Bond) UPH is "Per Head" — divide machines needed by heads/machine for actual count'],
    ['- Blank machine-hours cells mean the process does not apply to that package (UPH = 0)'],
    ['- DDR values are in K units (thousands)'],
    ['- Products with DDR = 0 are excluded'],
    [''],
    ['Mapping Notes:'],
    ['- 8LEIAJ → mapped to 8SOIJ (EIAJ is the SOIJ package standard)'],
    ['- 20LSSOP UD → mapped to 20SSOP-HD (upside-down die variant)'],
    ['- 8LSOIC IDF → mapped to 8SOIC (IDF is a die-flip variant)'],
    ['- 14LSOIC(Plasma) → mapped to 14SOIC (plasma cleaning variant)'],
    ['- Multiple DDR products may map to the same UPH package (see "By UPH Package" sheet)'],
]
for i, row in enumerate(notes):
    ws_note.cell(i + 1, 1, row[0] if row else '')
    if i == 0:
        ws_note.cell(1, 1).font = Font(name='Calibri', size=14, bold=True, color='1B3A5C')
    elif row and row[0].startswith('IMPORTANT'):
        ws_note.cell(i + 1, 1).font = Font(name='Calibri', size=11, bold=True, color='C0392B')
    elif row and row[0].startswith('HOW TO USE'):
        ws_note.cell(i + 1, 1).font = Font(name='Calibri', size=11, bold=True, color='2D8E4E')
    else:
        ws_note.cell(i + 1, 1).font = Font(name='Calibri', size=11)
ws_note.column_dimensions['A'].width = 95

# ── Save ──
out_path = 'Output/Machine_Usage_vs_DDR.xlsx'
wb_out.save(out_path)
print(f'Saved to {out_path}')
print(f'Total DDR products: {N}')
print(f'Unique UPH packages: {len(sorted_pkgs)}')
print(f'Sheets: Machine Hours Detail | Summary by Process | By UPH Package | Reverse Calc | UPH Ref (hidden) | Notes')
print(f'\nAll calculations are Excel FORMULAS — edit "Future DDR" column (D) to auto-recalculate.')
