import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(r"D:\claude\Project\raw\BOI 15 upgrade up Mar 20'26_finance confirm.xlsx", data_only=True)
ws = wb['Upgrade BOI10&11 to 15']

print('='*80)
print('ASSEMBLY UPGRADE - Items with PR / Spending')
print('='*80)
current_process = ''
assy_items = []

for r in range(7, 79):  # Assembly section rows 7-78
    a = ws.cell(r,1).value
    b = ws.cell(r,2).value
    if not b:
        continue
    if a and isinstance(a, str):
        current_process = a.strip()

    pr = ws.cell(r,17).value
    spent = ws.cell(r,18).value
    remain = ws.cell(r,19).value
    total_usd = ws.cell(r,12).value or 0
    plant = ws.cell(r,14).value
    remark = ws.cell(r,20).value

    item = {
        'row': r, 'process': current_process,
        'detail': str(b)[:70], 'plant': plant,
        'total_usd': total_usd, 'pr': pr,
        'spent': float(spent) if spent else 0,
        'remain': remain, 'remark': remark
    }
    assy_items.append(item)

# Print items with spending
print('\n--- Items with PR/Spending ---')
for i in assy_items:
    if i['pr'] or i['spent'] > 0:
        print(f"  {i['process']:22s} | {i['plant']:5s} | {i['detail']:70s}")
        print(f"  {'':22s} | Plan: ${i['total_usd']:>10,.2f} | PR: {i['pr']} | Spent: ${i['spent']:>10,.2f} | Remain: {i['remain']}")

print(f"\nTotal ASSY upgrade items: {len(assy_items)}")
total_plan = sum(i['total_usd'] for i in assy_items)
total_spent = sum(i['spent'] for i in assy_items)
print(f"ASSY Upgrade Total Plan: ${total_plan:,.2f}")
print(f"ASSY Upgrade Total Spent: ${total_spent:,.2f}")

for plant in ['MTAI', 'MMT']:
    items = [i for i in assy_items if i['plant'] == plant]
    plan = sum(i['total_usd'] for i in items)
    spent = sum(i['spent'] for i in items)
    pr_items = [i for i in items if i['pr']]
    print(f"  {plant}: Plan ${plan:,.2f} | Spent ${spent:,.2f} | Items: {len(items)} | With PR: {len(pr_items)}")

# NEW INVESTMENT analysis from PR sheet
print('\n' + '='*80)
print('NEW INVESTMENT PRs (from PR sheet)')
print('='*80)
ws_pr = wb['PR']
new_invest_prs = []
for r in range(3, ws_pr.max_row+1):
    pr_num = ws_pr.cell(r, 4).value
    pr_usd = ws_pr.cell(r, 5).value
    if pr_num and pr_usd:
        new_invest_prs.append({'pr': pr_num, 'usd': pr_usd})
        print(f"  PR# {pr_num} | ${pr_usd:>12,.2f}")

total_new = sum(p['usd'] for p in new_invest_prs)
print(f"\nTotal New Investment PRs: {len(new_invest_prs)} | Total: ${total_new:,.2f}")

# Summary from Sum 1M
print('\n' + '='*80)
print('ASSEMBLY BUDGET SUMMARY (from Sum 1M sheet)')
print('='*80)
ws_sum = wb['Sum 1M']
for r in [6, 7]:  # Assembly MTHAI and MMT rows
    name = ws_sum.cell(r, 1).value
    upg_plan = ws_sum.cell(r, 2).value or 0
    upg_actual = ws_sum.cell(r, 3).value or 0
    upg_remain_k = ws_sum.cell(r, 5).value or 0
    ni_plan = ws_sum.cell(r, 6).value or 0
    ni_actual = ws_sum.cell(r, 7).value or 0
    ni_remain_k = ws_sum.cell(r, 9).value or 0
    grand_plan = ws_sum.cell(r, 10).value or 0
    grand_actual_k = ws_sum.cell(r, 11).value or 0
    grand_remain_k = ws_sum.cell(r, 12).value or 0
    print(f"\n{name}:")
    print(f"  Upgrade:    Plan ${upg_plan}K | Actual ${upg_actual:,.2f} | Remain ${upg_remain_k:.2f}K")
    print(f"  New Invest: Plan ${ni_plan}K | Actual ${ni_actual:,.2f} | Remain ${ni_remain_k:.2f}K")
    print(f"  Grand:      Plan ${grand_plan}K | Actual ${grand_actual_k:.2f}K | Remain ${grand_remain_k:.2f}K")

# New investment mapping from Summary sheet
print('\n' + '='*80)
print('NEW INVESTMENT BY DEPARTMENT (from Summary sheet)')
print('='*80)
ws_summary = wb["Summary(P'TAN)"]
for r in range(5, 14):
    dept = ws_summary.cell(r, 7).value
    actual = ws_summary.cell(r, 8).value
    if dept:
        print(f"  {str(dept):15s} | Actual: ${actual:>12,.2f}" if actual else f"  {str(dept):15s} | Actual: $0")
