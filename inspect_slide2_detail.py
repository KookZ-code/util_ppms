# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
import lxml.etree as etree

EMU_PER_INCH = 914400

def emu_to_in(emu):
    if emu is None: return None
    return round(emu / EMU_PER_INCH, 4)

def get_rgb(color_obj):
    try:
        if color_obj.type is not None:
            r = color_obj.rgb
            return f"#{r.r:02X}{r.g:02X}{r.b:02X}"
    except: pass
    return "(theme/none)"

pptx_path = r"D:\claude\Project\Output\DA_ASSY_CAP_Exec_Summary_Apr22.pptx"
prs = Presentation(pptx_path)
slide = prs.slides[1]  # Slide 2 (0-indexed)

print("=" * 80)
print("SLIDE 2 — DETAILED INSPECTION")
print("=" * 80)
print(f"Number of shapes: {len(slide.shapes)}")

for si, shape in enumerate(slide.shapes):
    print(f"\n{'─'*60}")
    print(f"Shape #{si+1}: '{shape.name}'")
    print(f"  Type: {shape.shape_type}")
    print(f"  Left:   {emu_to_in(shape.left)}\"  ({shape.left} EMU)")
    print(f"  Top:    {emu_to_in(shape.top)}\"  ({shape.top} EMU)")
    print(f"  Width:  {emu_to_in(shape.width)}\"  ({shape.width} EMU)")
    print(f"  Height: {emu_to_in(shape.height)}\"  ({shape.height} EMU)")
    print(f"  Right edge:  {emu_to_in(shape.left + shape.width)}\"")
    print(f"  Bottom edge: {emu_to_in(shape.top + shape.height)}\"")

    # Fill
    try:
        fill = shape.fill
        print(f"  Fill.type: {fill.type}")
        try:
            print(f"  Fill.fore_color: {get_rgb(fill.fore_color)}")
        except: pass
    except: pass

    # Line
    try:
        line = shape.line
        lw = line.width
        lc = None
        try: lc = get_rgb(line.color)
        except: pass
        print(f"  Line: width={lw}, color={lc}")
    except: pass

    # Text
    if shape.has_text_frame:
        tf = shape.text_frame
        print(f"  TextFrame:")
        print(f"    word_wrap={tf.word_wrap}")
        print(f"    auto_size={tf.auto_size}")
        try: print(f"    margin_left={emu_to_in(tf.margin_left)}\"")
        except: pass
        try: print(f"    margin_top={emu_to_in(tf.margin_top)}\"")
        except: pass
        try: print(f"    margin_right={emu_to_in(tf.margin_right)}\"")
        except: pass
        try: print(f"    margin_bottom={emu_to_in(tf.margin_bottom)}\"")
        except: pass

        for pi, para in enumerate(tf.paragraphs):
            text = para.text
            if not text.strip() and pi > 0:
                print(f"    Para[{pi}]: (empty)")
                continue
            print(f"    Para[{pi}]: '{text}'")

            # Paragraph-level font (pPr/defRPr)
            pPr = para._pPr
            if pPr is not None:
                defRPr = pPr.find(qn('a:defRPr'))
                if defRPr is not None:
                    sz = defRPr.get('sz')
                    b = defRPr.get('b')
                    print(f"      para defRPr: sz={sz}, bold={b}")

            print(f"      alignment={para.alignment}")
            print(f"      space_before={para.space_before}")
            print(f"      space_after={para.space_after}")

            for ri, run in enumerate(para.runs):
                f = run.font
                fc = None
                try: fc = get_rgb(f.color)
                except: pass
                sz = f.size.pt if f.size else None
                print(f"      Run[{ri}]: '{run.text}'  size={sz}pt bold={f.bold} italic={f.italic} name='{f.name}' color={fc}")

    # Table details
    if shape.has_table:
        tbl = shape.table
        print(f"  Table: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
        for r_idx, row in enumerate(tbl.rows):
            print(f"  Row {r_idx}: height={emu_to_in(row.height)}\"")
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text_frame.text
                # Get cell fill
                try:
                    cf = cell.fill
                    cf_type = cf.type
                    cf_color = None
                    try: cf_color = get_rgb(cf.fore_color)
                    except: pass
                except:
                    cf_type = None
                    cf_color = None

                # Get cell font
                tf = cell.text_frame
                cell_font_size = None
                cell_font_bold = None
                cell_font_color = None
                cell_font_name = None
                if tf.paragraphs and tf.paragraphs[0].runs:
                    run = tf.paragraphs[0].runs[0]
                    f = run.font
                    cell_font_size = f.size.pt if f.size else None
                    cell_font_bold = f.bold
                    try: cell_font_color = get_rgb(f.color)
                    except: pass
                    cell_font_name = f.name

                print(f"    Cell[{r_idx},{c_idx}]: '{cell_text}'")
                print(f"      fill_type={cf_type}, fill_color={cf_color}")
                print(f"      font: size={cell_font_size}pt, bold={cell_font_bold}, color={cell_font_color}, name={cell_font_name}")

print("\n\nDONE.")
