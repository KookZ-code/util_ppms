import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

def norm_prefix(p):
    return re.sub(r'[^A-Z0-9]', '', p.upper())

def parse_mc_no(mc_str):
    """Parse MC No. field into a set of (normalized_prefix, number) tuples."""
    if not mc_str:
        return set()
    machines = set()
    current_prefix = ''
    s = str(mc_str).strip()
    # Strip parenthetical content (serial numbers, notes) e.g., "TRY-09 (11053011-L)" -> "TRY-09"
    s = re.sub(r'\([^)]*\)', '', s)
    # Replace spaces-as-separator between tokens like "TRY#004 TRY#005" -> "TRY#004,TRY#005"
    s = re.sub(r'(\d)\s+([A-Z])', r'\1,\2', s)
    s = re.sub(r'(\d)\s*([A-Z/])', r'\1,\2', s)
    segments = [seg.strip() for seg in s.split(',') if seg.strip()]
    for seg in segments:
        tokens = re.findall(r'([^\d,]*?)(\d+)', seg)
        if not tokens:
            continue
        is_range = False
        if len(tokens) == 2:
            first_num_end = seg.index(tokens[0][1]) + len(tokens[0][1])
            rest = seg[first_num_end:]
            if '-' in rest and re.search(r'\d', rest):
                is_range = True
        if is_range:
            prefix1 = tokens[0][0].strip()
            num1 = int(tokens[0][1])
            num2 = int(tokens[1][1])
            if prefix1:
                current_prefix = norm_prefix(prefix1)
            for n in range(min(num1, num2), max(num1, num2) + 1):
                machines.add((current_prefix, n))
        else:
            for pfx, num in tokens:
                pfx = pfx.strip()
                if pfx:
                    current_prefix = norm_prefix(pfx)
                machines.add((current_prefix, int(num)))
    return machines

SRC = r"D:\claude\Project\raw\BOI 15 upgrade up Mar 20'26_finance confirm.xlsx"
OUT = r"D:\claude\Project\Output\BOI15_Assembly_Budget_Analysis.xlsx"

wb_src = openpyxl.load_workbook(SRC, data_only=True)

# ============ STYLES ============
DARK_BLUE = "1B3A5C"
MED_BLUE = "2E75B6"
GREEN = "2D8E4E"
ORANGE = "E67E22"
LIGHT_BLUE = "E8F0FE"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4E1"
LIGHT_ORANGE = "FFF2CC"
WHITE = "FFFFFF"

hdr_font = Font(name="Arial", bold=True, color=WHITE, size=11)
hdr_fill = PatternFill("solid", fgColor=DARK_BLUE)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
data_font = Font(name="Arial", size=10)
title_font = Font(name="Arial", bold=True, size=16, color=DARK_BLUE)
section_font = Font(name="Arial", bold=True, size=13, color=DARK_BLUE)
thin_border = Border(
    left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0'))
alt_fill_1 = PatternFill("solid", fgColor=LIGHT_BLUE)
alt_fill_2 = PatternFill("solid", fgColor=WHITE)

def style_header_row(ws, row, cols, font=hdr_font, fill=hdr_fill):
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = font
        cell.fill = fill
        cell.alignment = hdr_align
        cell.border = thin_border

def style_data_row(ws, row, cols, idx=0):
    fill = alt_fill_1 if idx % 2 == 0 else alt_fill_2
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = data_font
        cell.border = thin_border
        cell.fill = fill

def style_total_row(ws, row, cols, fill_color=MED_BLUE):
    for c in range(1, cols+1):
        cell = ws.cell(row, c)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.border = thin_border

# ============ PARSE SOURCE DATA ============
ws_upg = wb_src['Upgrade BOI10&11 to 15']
ws_pr = wb_src['PR']
ws_sum1m = wb_src['Sum 1M']
ws_summary = wb_src["Summary(P'TAN)"]

# --- Process -> Machine type mapping ---
process_map = {
    'SAW': 'SAW', 'SAW QFN': 'SAW', 'SAW(MMT)': 'SAW',
    'DIE ATTACH': 'DIEBONDER', 'DIE ATTACH (MMT)': 'DIEBONDER',
    'WIRE BONDING': 'Wirebonder', 'WIRE BONDING (MMT)': 'Wirebonder',
    'Module (SPG)': 'Module - Reflow', 'Final leak (SPG)': 'Module - Reflow',
    'Backgrind': 'BACKGRIND', 'Mounter': 'WAFERMOUNT',
    'MOLD': 'MOLD', 'Plating': 'SOLDERPLATE',
    'MARK': 'LASERMARK', 'TRIM/FORM': 'TRIMFORM',
    'Form/Sing': 'TRIMFORM', 'ISOLATE': 'ISOLATE',
}

# --- Assembly Upgrade items with machine type ---
assy_upgrade = []
current_process = ''
mch_data = defaultdict(lambda: {'plan': 0, 'spent': 0, 'remain': 0, 'count': 0, 'pr_count': 0, 'items': [], 'machines': set()})

for r in range(7, 79):
    b = ws_upg.cell(r, 2).value
    if not b:
        continue
    a = ws_upg.cell(r, 1).value
    if a and isinstance(a, str):
        current_process = a.strip()

    plant = ws_upg.cell(r, 14).value or ''
    total_usd = ws_upg.cell(r, 12).value or 0
    spent = float(ws_upg.cell(r, 18).value or 0)
    remain = float(ws_upg.cell(r, 19).value or 0) if ws_upg.cell(r, 19).value else 0
    pr = ws_upg.cell(r, 17).value
    detail = str(b).strip()

    # Determine machine type
    mch = process_map.get(current_process, current_process)
    dl = detail.lower()
    if 'x-ray' in dl or 'xray' in dl or 'x ray' in dl:
        mch = 'XRAY'
    elif 'plasma' in dl or 'rf generator' in dl:
        mch = 'PLASMA'
    elif ('die shear' in dl or 'bond shear' in dl or 'wire pull' in dl) and mch != 'Wirebonder':
        mch = 'DIE SHEAR TESTER'
    elif 'microscope' in dl:
        mch = 'QA'

    item = {
        'process': current_process, 'mch_type': mch, 'detail': detail,
        'propose': ws_upg.cell(r, 3).value,
        'mc_no': ws_upg.cell(r, 6).value,
        'qty_set': ws_upg.cell(r, 7).value,
        'qty_mc': ws_upg.cell(r, 8).value,
        'total_usd': total_usd, 'plant': plant,
        'boi': ws_upg.cell(r, 15).value,
        'contact': ws_upg.cell(r, 13).value,
        'pr': pr, 'spent': spent, 'remain': remain,
        'remark': ws_upg.cell(r, 20).value,
    }
    assy_upgrade.append(item)

    key = (plant, mch)
    mch_data[key]['plan'] += total_usd
    mch_data[key]['spent'] += spent
    mch_data[key]['remain'] += remain
    mch_data[key]['count'] += 1
    if pr:
        mch_data[key]['pr_count'] += 1
    mch_data[key]['machines'] |= parse_mc_no(item['mc_no'])

# --- New Investment PRs ---
new_invest_prs = []
for r in range(3, ws_pr.max_row + 1):
    pr_num = ws_pr.cell(r, 4).value
    pr_usd = ws_pr.cell(r, 5).value
    if pr_num and pr_usd:
        new_invest_prs.append({'pr': str(pr_num), 'usd': pr_usd})

# --- Upgrade PRs ---
upgrade_prs = []
for r in range(3, ws_pr.max_row + 1):
    pr_num = ws_pr.cell(r, 1).value
    pr_usd = ws_pr.cell(r, 2).value
    if pr_num and pr_usd:
        upgrade_prs.append({'pr': str(pr_num), 'usd': pr_usd})

# --- NI depts ---
ni_depts = {}
for r in range(5, 14):
    dept = ws_summary.cell(r, 7).value
    actual = ws_summary.cell(r, 8).value or 0
    if dept:
        ni_depts[dept.strip()] = actual

# --- Sum 1M ---
sum1m = {}
for r in [6, 7]:
    name = ws_sum1m.cell(r, 1).value
    sum1m[name] = {
        'upg_plan_k': ws_sum1m.cell(r, 2).value or 0,
        'upg_actual': ws_sum1m.cell(r, 3).value or 0,
        'ni_plan_k': ws_sum1m.cell(r, 6).value or 0,
        'ni_actual': ws_sum1m.cell(r, 7).value or 0,
        'ni_remain_k': ws_sum1m.cell(r, 9).value or 0,
    }

# ============ CREATE WORKBOOK ============
wb = Workbook()

# ============================
# SHEET 1: DASHBOARD
# ============================
ws1 = wb.active
ws1.title = "Assembly Dashboard"
ws1.sheet_properties.tabColor = DARK_BLUE

ws1.cell(1, 1, "BOI 15 - Assembly Budget Analysis").font = Font(name="Arial", bold=True, size=18, color=DARK_BLUE)
ws1.merge_cells("A1:H1")
ws1.cell(2, 1, "Upgrade + New Investment | MTAI & MMT | As of Mar 2026 (Finance Confirmed)").font = Font(name="Arial", size=11, color="666666")
ws1.merge_cells("A2:H2")

# --- Budget Overview ---
r = 4
ws1.cell(r, 1, "ASSEMBLY BUDGET OVERVIEW ($K)").font = section_font
ws1.merge_cells(f"A{r}:G{r}")

r = 6
for c, h in enumerate(["Plant", "Type", "Budget Plan ($K)", "Actual Spent ($K)", "Remain ($K)", "% Used", "Status"], 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 7)

mtai = sum1m['Assembly MTHAI']
mmt = sum1m['Assembly MMT']

budget_rows = [
    ("MTAI", "Upgrade", mtai['upg_plan_k'], mtai['upg_actual'] / 1000),
    ("MTAI", "New Invest", mtai['ni_plan_k'], mtai['ni_actual'] / 1000),
    None,  # subtotal
    ("MMT", "Upgrade", mmt['upg_plan_k'], mmt['upg_actual'] / 1000),
    ("MMT", "New Invest", mmt['ni_plan_k'], mmt['ni_actual'] / 1000),
    None,  # subtotal
]

r = 7
idx = 0
for item in budget_rows:
    if item is None:
        # Subtotal
        ws1.cell(r, 1, f"{ws1.cell(r-2,1).value} Total")
        ws1.cell(r, 3).value = f'=C{r-2}+C{r-1}'
        ws1.cell(r, 4).value = f'=D{r-2}+D{r-1}'
        ws1.cell(r, 5).value = f'=C{r}-D{r}'
        ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
        ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
        style_total_row(ws1, r, 7, MED_BLUE)
    else:
        plant, typ, plan, actual = item
        ws1.cell(r, 1, plant)
        ws1.cell(r, 2, typ)
        ws1.cell(r, 3, plan)
        ws1.cell(r, 4, actual)
        ws1.cell(r, 5).value = f'=C{r}-D{r}'
        ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
        ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
        style_data_row(ws1, r, 7, idx)
        idx += 1
    for c in [3, 4, 5]:
        ws1.cell(r, c).number_format = '#,##0.00'
    ws1.cell(r, 6).number_format = '0.0%'
    r += 1

# Grand total
r += 0  # r is already at next row
ws1.cell(r, 1, "ASSEMBLY GRAND TOTAL")
ws1.cell(r, 3).value = f'=C9+C12'
ws1.cell(r, 4).value = f'=D9+D12'
ws1.cell(r, 5).value = f'=C{r}-D{r}'
ws1.cell(r, 6).value = f'=IF(C{r}=0,"",D{r}/C{r})'
ws1.cell(r, 7).value = f'=IF(F{r}>=1,"Over Budget",IF(F{r}>=0.8,"Near Limit","On Track"))'
style_total_row(ws1, r, 7, DARK_BLUE)
for c in [3, 4, 5]:
    ws1.cell(r, c).number_format = '#,##0.00'
ws1.cell(r, 6).number_format = '0.0%'
grand_row = r

# ======================================================
# UPGRADE BY MACHINE TYPE - MTAI
# ======================================================
r = grand_row + 3
ws1.cell(r, 1, "UPGRADE BY MACHINE TYPE - MTAI").font = section_font
ws1.merge_cells(f"A{r}:I{r}")

r += 2
mch_headers = ["Machine Type", "Items", "Unique M/C", "Plan (USD)", "Spent (USD)", "Pending (USD)", "% Spent", "PRs Issued", "Key Items"]
for c, h in enumerate(mch_headers, 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 9)
mtai_hdr_row = r

r += 1
mtai_first = r
mtai_keys = sorted([k for k in mch_data if k[0] == 'MTAI'], key=lambda x: x[1])
mtai_total_machines = set()
for idx, key in enumerate(mtai_keys):
    v = mch_data[key]
    ws1.cell(r, 1, key[1])
    ws1.cell(r, 2, v['count'])
    ws1.cell(r, 3, len(v['machines']))
    ws1.cell(r, 4, v['plan'])
    ws1.cell(r, 5, v['spent'])
    ws1.cell(r, 6, v['plan'] - v['spent'])
    ws1.cell(r, 7).value = f'=IF(D{r}=0,"",E{r}/D{r})'
    ws1.cell(r, 8, v['pr_count'])
    key_items = [i['detail'][:45] for i in assy_upgrade if i['plant'] == 'MTAI' and i['mch_type'] == key[1] and i['spent'] > 0]
    ws1.cell(r, 9, '; '.join(key_items) if key_items else '')
    style_data_row(ws1, r, 9, idx)
    for c in [4, 5, 6]: ws1.cell(r, c).number_format = '#,##0.00'
    ws1.cell(r, 7).number_format = '0.0%'
    for c in [2, 3, 8]: ws1.cell(r, c).alignment = Alignment(horizontal="center")
    mtai_total_machines |= v['machines']
    r += 1

# MTAI subtotal
ws1.cell(r, 1, "MTAI TOTAL")
ws1.cell(r, 2).value = f'=SUM(B{mtai_first}:B{r-1})'
ws1.cell(r, 3, len(mtai_total_machines))
ws1.cell(r, 4).value = f'=SUM(D{mtai_first}:D{r-1})'
ws1.cell(r, 5).value = f'=SUM(E{mtai_first}:E{r-1})'
ws1.cell(r, 6).value = f'=SUM(F{mtai_first}:F{r-1})'
ws1.cell(r, 7).value = f'=IF(D{r}=0,"",E{r}/D{r})'
ws1.cell(r, 8).value = f'=SUM(H{mtai_first}:H{r-1})'
style_total_row(ws1, r, 9, MED_BLUE)
for c in [4, 5, 6]: ws1.cell(r, c).number_format = '#,##0.00'
ws1.cell(r, 7).number_format = '0.0%'
ws1.cell(r, 3).alignment = Alignment(horizontal="center")
mtai_total_row = r

# ======================================================
# UPGRADE BY MACHINE TYPE - MMT
# ======================================================
r += 2
ws1.cell(r, 1, "UPGRADE BY MACHINE TYPE - MMT").font = section_font
ws1.merge_cells(f"A{r}:I{r}")

r += 2
for c, h in enumerate(mch_headers, 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 9)

r += 1
mmt_first = r
mmt_keys = sorted([k for k in mch_data if k[0] == 'MMT'], key=lambda x: x[1])
mmt_total_machines = set()
for idx, key in enumerate(mmt_keys):
    v = mch_data[key]
    ws1.cell(r, 1, key[1])
    ws1.cell(r, 2, v['count'])
    ws1.cell(r, 3, len(v['machines']))
    ws1.cell(r, 4, v['plan'])
    ws1.cell(r, 5, v['spent'])
    ws1.cell(r, 6, v['plan'] - v['spent'])
    ws1.cell(r, 7).value = f'=IF(D{r}=0,"",E{r}/D{r})'
    ws1.cell(r, 8, v['pr_count'])
    key_items = [i['detail'][:45] for i in assy_upgrade if i['plant'] == 'MMT' and i['mch_type'] == key[1] and i['spent'] > 0]
    ws1.cell(r, 9, '; '.join(key_items) if key_items else '')
    style_data_row(ws1, r, 9, idx)
    for c in [4, 5, 6]: ws1.cell(r, c).number_format = '#,##0.00'
    ws1.cell(r, 7).number_format = '0.0%'
    for c in [2, 3, 8]: ws1.cell(r, c).alignment = Alignment(horizontal="center")
    mmt_total_machines |= v['machines']
    r += 1

# MMT subtotal
ws1.cell(r, 1, "MMT TOTAL")
ws1.cell(r, 2).value = f'=SUM(B{mmt_first}:B{r-1})'
ws1.cell(r, 3, len(mmt_total_machines))
ws1.cell(r, 4).value = f'=SUM(D{mmt_first}:D{r-1})'
ws1.cell(r, 5).value = f'=SUM(E{mmt_first}:E{r-1})'
ws1.cell(r, 6).value = f'=SUM(F{mmt_first}:F{r-1})'
ws1.cell(r, 7).value = f'=IF(D{r}=0,"",E{r}/D{r})'
ws1.cell(r, 8).value = f'=SUM(H{mmt_first}:H{r-1})'
style_total_row(ws1, r, 9, MED_BLUE)
for c in [4, 5, 6]: ws1.cell(r, c).number_format = '#,##0.00'
ws1.cell(r, 7).number_format = '0.0%'
ws1.cell(r, 3).alignment = Alignment(horizontal="center")
mmt_total_row = r

# ASSY Grand total
r += 1
ws1.cell(r, 1, "ASSEMBLY UPGRADE GRAND TOTAL")
ws1.cell(r, 2).value = f'=B{mtai_total_row}+B{mmt_total_row}'
ws1.cell(r, 3).value = f'=C{mtai_total_row}+C{mmt_total_row}'
ws1.cell(r, 4).value = f'=D{mtai_total_row}+D{mmt_total_row}'
ws1.cell(r, 5).value = f'=E{mtai_total_row}+E{mmt_total_row}'
ws1.cell(r, 6).value = f'=F{mtai_total_row}+F{mmt_total_row}'
ws1.cell(r, 7).value = f'=IF(D{r}=0,"",E{r}/D{r})'
ws1.cell(r, 8).value = f'=H{mtai_total_row}+H{mmt_total_row}'
style_total_row(ws1, r, 9, DARK_BLUE)
for c in [4, 5, 6]: ws1.cell(r, c).number_format = '#,##0.00'
ws1.cell(r, 7).number_format = '0.0%'

# --- New Investment summary ---
r += 3
ws1.cell(r, 1, "ASSEMBLY NEW INVESTMENT (Summary)").font = section_font
ws1.merge_cells(f"A{r}:E{r}")

r += 2
for c, h in enumerate(["Department", "Budget Plan ($K)", "Actual (USD)", "Remain ($K)"], 1):
    ws1.cell(r, c, h)
style_header_row(ws1, r, 4)

r += 1
ni_first = r
ws1.cell(r, 1, "ASSY MTHAI")
ws1.cell(r, 2, mtai['ni_plan_k'])
ws1.cell(r, 3, ni_depts.get('ASSY MTHAI', 0))
ws1.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws1, r, 4, 0)
ws1.cell(r, 3).number_format = '#,##0.00'
for c in [2, 4]: ws1.cell(r, c).number_format = '#,##0.00'

r += 1
ws1.cell(r, 1, "ASSY MMT")
ws1.cell(r, 2, mmt['ni_plan_k'])
ws1.cell(r, 3, ni_depts.get('ASSY MMT', 0))
ws1.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws1, r, 4, 1)
ws1.cell(r, 3).number_format = '#,##0.00'
for c in [2, 4]: ws1.cell(r, c).number_format = '#,##0.00'

r += 1
ws1.cell(r, 1, "ASSY NI TOTAL")
ws1.cell(r, 2).value = f'=B{ni_first}+B{ni_first+1}'
ws1.cell(r, 3).value = f'=C{ni_first}+C{ni_first+1}'
ws1.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_total_row(ws1, r, 4, DARK_BLUE)
for c in [2, 3, 4]: ws1.cell(r, c).number_format = '#,##0.00'

# Column widths
for c, w in {1:28, 2:10, 3:12, 4:16, 5:16, 6:16, 7:10, 8:12, 9:50}.items():
    ws1.column_dimensions[get_column_letter(c)].width = w
ws1.freeze_panes = "A3"

# ============================
# SHEET 2: MTAI UPGRADE DETAIL
# ============================
ws_mtai = wb.create_sheet("MTAI Upgrade Detail")
ws_mtai.sheet_properties.tabColor = MED_BLUE

ws_mtai.cell(1, 1, "MTAI - Assembly Upgrade Items Detail").font = title_font
ws_mtai.merge_cells("A1:P1")

detail_headers = ["No.", "Machine Type", "Process", "Upgrade Details", "MC No.", "Qty/Set", "Qty MC",
                   "Total (USD)", "Contact", "BOI", "PR#", "Spent (USD)", "Remain (USD)", "% Spent", "Status", "Remark"]
r = 3
for c, h in enumerate(detail_headers, 1):
    ws_mtai.cell(r, c, h)
style_header_row(ws_mtai, r, 16)

r = 4
mtai_items = sorted([i for i in assy_upgrade if i['plant'] == 'MTAI'],
                    key=lambda x: (0 if x['spent'] > 0 else (1 if x['pr'] else 2), x['mch_type']))

for idx, item in enumerate(mtai_items):
    spent = item['spent']
    has_pr = item['pr'] is not None and str(item['pr']).strip() != ''
    if spent > 0 and (item['remain'] <= 0):
        status = "Completed"
    elif spent > 0:
        status = "In Progress"
    elif has_pr:
        status = "PR Issued"
    else:
        status = "Pending"

    ws_mtai.cell(r, 1, idx + 1)
    ws_mtai.cell(r, 2, item['mch_type'])
    ws_mtai.cell(r, 3, item['process'])
    ws_mtai.cell(r, 4, item['detail'])
    ws_mtai.cell(r, 5, item['mc_no'])
    ws_mtai.cell(r, 6, item['qty_set'])
    ws_mtai.cell(r, 7, item['qty_mc'])
    ws_mtai.cell(r, 8, item['total_usd'])
    ws_mtai.cell(r, 9, item['contact'])
    ws_mtai.cell(r, 10, str(item['boi']) if item['boi'] else '')
    ws_mtai.cell(r, 11, item['pr'])
    ws_mtai.cell(r, 12, item['spent'] if item['spent'] else None)
    ws_mtai.cell(r, 13, item['remain'] if item['remain'] else None)
    ws_mtai.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
    ws_mtai.cell(r, 15, status)
    ws_mtai.cell(r, 16, item['remark'])

    style_data_row(ws_mtai, r, 16, idx)
    ws_mtai.cell(r, 8).number_format = '#,##0.00'
    ws_mtai.cell(r, 12).number_format = '#,##0.00'
    ws_mtai.cell(r, 13).number_format = '#,##0.00'
    ws_mtai.cell(r, 14).number_format = '0.0%'
    for cc in [1, 2, 6, 7, 10, 15]:
        ws_mtai.cell(r, cc).alignment = Alignment(horizontal="center")

    if status == "Completed":
        ws_mtai.cell(r, 15).fill = PatternFill("solid", fgColor="C6EFCE")
        ws_mtai.cell(r, 15).font = Font(name="Arial", size=10, color="006100")
    elif status == "In Progress":
        ws_mtai.cell(r, 15).fill = PatternFill("solid", fgColor="FFEB9C")
        ws_mtai.cell(r, 15).font = Font(name="Arial", size=10, color="9C6500")
    elif status == "PR Issued":
        ws_mtai.cell(r, 15).fill = PatternFill("solid", fgColor="FCD5B4")
        ws_mtai.cell(r, 15).font = Font(name="Arial", size=10, color="974706")
    r += 1

last_mtai = r - 1
ws_mtai.cell(r, 1, "MTAI TOTAL")
ws_mtai.cell(r, 7).value = f'=SUM(G4:G{last_mtai})'
ws_mtai.cell(r, 8).value = f'=SUM(H4:H{last_mtai})'
ws_mtai.cell(r, 12).value = f'=SUM(L4:L{last_mtai})'
ws_mtai.cell(r, 13).value = f'=SUM(M4:M{last_mtai})'
ws_mtai.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
for c in [8, 12, 13]: ws_mtai.cell(r, c).number_format = '#,##0.00'
ws_mtai.cell(r, 14).number_format = '0.0%'
style_total_row(ws_mtai, r, 16, DARK_BLUE)

ws_mtai.auto_filter.ref = f"A3:P{last_mtai}"
ws_mtai.freeze_panes = "A4"
for c, w in {1:5, 2:18, 3:22, 4:55, 5:35, 6:8, 7:8, 8:13, 9:20, 10:8, 11:18, 12:13, 13:13, 14:9, 15:13, 16:30}.items():
    ws_mtai.column_dimensions[get_column_letter(c)].width = w

# ============================
# SHEET 3: MMT UPGRADE DETAIL
# ============================
ws_mmt = wb.create_sheet("MMT Upgrade Detail")
ws_mmt.sheet_properties.tabColor = GREEN

ws_mmt.cell(1, 1, "MMT - Assembly Upgrade Items Detail").font = title_font
ws_mmt.merge_cells("A1:P1")

r = 3
for c, h in enumerate(detail_headers, 1):
    ws_mmt.cell(r, c, h)
style_header_row(ws_mmt, r, 16)

r = 4
mmt_items = sorted([i for i in assy_upgrade if i['plant'] == 'MMT'],
                   key=lambda x: (0 if x['spent'] > 0 else (1 if x['pr'] else 2), x['mch_type']))

for idx, item in enumerate(mmt_items):
    spent = item['spent']
    has_pr = item['pr'] is not None and str(item['pr']).strip() != ''
    if spent > 0 and (item['remain'] <= 0):
        status = "Completed"
    elif spent > 0:
        status = "In Progress"
    elif has_pr:
        status = "PR Issued"
    else:
        status = "Pending"

    ws_mmt.cell(r, 1, idx + 1)
    ws_mmt.cell(r, 2, item['mch_type'])
    ws_mmt.cell(r, 3, item['process'])
    ws_mmt.cell(r, 4, item['detail'])
    ws_mmt.cell(r, 5, item['mc_no'])
    ws_mmt.cell(r, 6, item['qty_set'])
    ws_mmt.cell(r, 7, item['qty_mc'])
    ws_mmt.cell(r, 8, item['total_usd'])
    ws_mmt.cell(r, 9, item['contact'])
    ws_mmt.cell(r, 10, str(item['boi']) if item['boi'] else '')
    ws_mmt.cell(r, 11, item['pr'])
    ws_mmt.cell(r, 12, item['spent'] if item['spent'] else None)
    ws_mmt.cell(r, 13, item['remain'] if item['remain'] else None)
    ws_mmt.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
    ws_mmt.cell(r, 15, status)
    ws_mmt.cell(r, 16, item['remark'])

    style_data_row(ws_mmt, r, 16, idx)
    ws_mmt.cell(r, 8).number_format = '#,##0.00'
    ws_mmt.cell(r, 12).number_format = '#,##0.00'
    ws_mmt.cell(r, 13).number_format = '#,##0.00'
    ws_mmt.cell(r, 14).number_format = '0.0%'
    for cc in [1, 2, 6, 7, 10, 15]:
        ws_mmt.cell(r, cc).alignment = Alignment(horizontal="center")

    if status == "Completed":
        ws_mmt.cell(r, 15).fill = PatternFill("solid", fgColor="C6EFCE")
        ws_mmt.cell(r, 15).font = Font(name="Arial", size=10, color="006100")
    elif status == "In Progress":
        ws_mmt.cell(r, 15).fill = PatternFill("solid", fgColor="FFEB9C")
        ws_mmt.cell(r, 15).font = Font(name="Arial", size=10, color="9C6500")
    elif status == "PR Issued":
        ws_mmt.cell(r, 15).fill = PatternFill("solid", fgColor="FCD5B4")
        ws_mmt.cell(r, 15).font = Font(name="Arial", size=10, color="974706")
    r += 1

last_mmt = r - 1
ws_mmt.cell(r, 1, "MMT TOTAL")
ws_mmt.cell(r, 7).value = f'=SUM(G4:G{last_mmt})'
ws_mmt.cell(r, 8).value = f'=SUM(H4:H{last_mmt})'
ws_mmt.cell(r, 12).value = f'=SUM(L4:L{last_mmt})'
ws_mmt.cell(r, 13).value = f'=SUM(M4:M{last_mmt})'
ws_mmt.cell(r, 14).value = f'=IF(H{r}=0,"",L{r}/H{r})'
for c in [8, 12, 13]: ws_mmt.cell(r, c).number_format = '#,##0.00'
ws_mmt.cell(r, 14).number_format = '0.0%'
style_total_row(ws_mmt, r, 16, DARK_BLUE)

ws_mmt.auto_filter.ref = f"A3:P{last_mmt}"
ws_mmt.freeze_panes = "A4"
for c, w in {1:5, 2:18, 3:22, 4:55, 5:35, 6:8, 7:8, 8:13, 9:20, 10:8, 11:18, 12:13, 13:13, 14:9, 15:13, 16:30}.items():
    ws_mmt.column_dimensions[get_column_letter(c)].width = w

# ============================
# SHEET 4: NEW INVESTMENT PRs
# ============================
ws3 = wb.create_sheet("New Investment PRs")
ws3.sheet_properties.tabColor = "2D8E4E"

ws3.cell(1, 1, "New Investment - PR List (All Departments)").font = title_font
ws3.merge_cells("A1:F1")
ws3.cell(2, 1, "Assembly items highlighted in green | Source: PR sheet").font = Font(name="Arial", size=10, color="666666")

r = 4
ws3.cell(r, 1, "Assembly New Investment Summary").font = section_font
ws3.merge_cells(f"A{r}:D{r}")

r = 6
for c, h in enumerate(["Department", "Budget Plan ($K)", "Actual (USD)", "Remain ($K)"], 1):
    ws3.cell(r, c, h)
style_header_row(ws3, r, 4)

r = 7
ws3.cell(r, 1, "ASSY MTHAI")
ws3.cell(r, 2, mtai['ni_plan_k'])
ws3.cell(r, 3, ni_depts.get('ASSY MTHAI', 0))
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws3, r, 4, 0)
ws3.cell(r, 3).number_format = '#,##0.00'
for c in [2, 4]: ws3.cell(r, c).number_format = '#,##0.00'

r = 8
ws3.cell(r, 1, "ASSY MMT")
ws3.cell(r, 2, mmt['ni_plan_k'])
ws3.cell(r, 3, ni_depts.get('ASSY MMT', 0))
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_data_row(ws3, r, 4, 1)
ws3.cell(r, 3).number_format = '#,##0.00'
for c in [2, 4]: ws3.cell(r, c).number_format = '#,##0.00'

r = 9
ws3.cell(r, 1, "ASSY NI TOTAL")
ws3.cell(r, 2).value = '=B7+B8'
ws3.cell(r, 3).value = '=C7+C8'
ws3.cell(r, 4).value = f'=B{r}-C{r}/1000'
style_total_row(ws3, r, 4, DARK_BLUE)
for c in [2, 3, 4]: ws3.cell(r, c).number_format = '#,##0.00'

r = 11
ws3.cell(r, 1, "New Investment - PR Detail List").font = section_font

r = 13
for c, h in enumerate(["No.", "PR#", "Amount (USD)", "Note"], 1):
    ws3.cell(r, c, h)
style_header_row(ws3, r, 4)

r = 14
assy_mthai_pr = '100104009'
for idx, p in enumerate(new_invest_prs):
    ws3.cell(r, 1, idx + 1)
    ws3.cell(r, 2, p['pr'])
    ws3.cell(r, 3, p['usd'])
    ws3.cell(r, 3).number_format = '#,##0.00'
    if p['pr'] == assy_mthai_pr:
        ws3.cell(r, 4, "ASSY MTHAI (confirmed)")
    style_data_row(ws3, r, 4, idx)
    if p['pr'] == assy_mthai_pr:
        for c in range(1, 5):
            ws3.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            ws3.cell(r, c).font = Font(name="Arial", size=10, bold=True, color="006100")
    r += 1

ws3.cell(r, 1, "TOTAL")
ws3.cell(r, 3).value = f'=SUM(C14:C{r-1})'
ws3.cell(r, 3).number_format = '#,##0.00'
style_total_row(ws3, r, 4, DARK_BLUE)

ws3.column_dimensions['A'].width = 6
ws3.column_dimensions['B'].width = 18
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 30
ws3.freeze_panes = "A3"

# ============================
# SHEET 5: ALL UPGRADE PRs
# ============================
ws4 = wb.create_sheet("Upgrade PRs")
ws4.sheet_properties.tabColor = ORANGE

ws4.cell(1, 1, "Upgrade - PR List (All Departments)").font = title_font
ws4.merge_cells("A1:E1")

pr_to_item = {}
for item in assy_upgrade:
    if item['pr']:
        for pr_raw in str(item['pr']).split(','):
            pr_clean = pr_raw.strip()
            if pr_clean:
                pr_to_item[pr_clean] = f"{item['plant']} | {item['mch_type']} - {str(item['detail'])[:40]}"

r = 3
for c, h in enumerate(["No.", "PR#", "Amount (USD)", "Linked Item (Plant | Machine - Detail)", "Type"], 1):
    ws4.cell(r, c, h)
style_header_row(ws4, r, 5)

r = 4
for idx, p in enumerate(upgrade_prs):
    ws4.cell(r, 1, idx + 1)
    ws4.cell(r, 2, p['pr'])
    ws4.cell(r, 3, p['usd'])
    ws4.cell(r, 3).number_format = '#,##0.00'
    linked = pr_to_item.get(p['pr'], '')
    ws4.cell(r, 4, linked)
    ws4.cell(r, 5, "ASSY" if linked else "TEST")
    style_data_row(ws4, r, 5, idx)
    if linked:
        for c in range(1, 6):
            ws4.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    r += 1

ws4.cell(r, 1, "TOTAL")
ws4.cell(r, 3).value = f'=SUM(C4:C{r-1})'
ws4.cell(r, 3).number_format = '#,##0.00'
style_total_row(ws4, r, 5, DARK_BLUE)

ws4.auto_filter.ref = f"A3:E{r-1}"
ws4.freeze_panes = "A4"
ws4.column_dimensions['A'].width = 6
ws4.column_dimensions['B'].width = 18
ws4.column_dimensions['C'].width = 18
ws4.column_dimensions['D'].width = 55
ws4.column_dimensions['E'].width = 10

# ============ SAVE ============
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Sheets: {wb.sheetnames}")
print(f"MTAI machine types: {len(mtai_keys)} | items: {len(mtai_items)}")
print(f"MMT machine types: {len(mmt_keys)} | items: {len(mmt_items)}")
