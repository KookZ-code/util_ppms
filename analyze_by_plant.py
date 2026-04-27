import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook(r"D:\claude\Project\raw\BOI 15 upgrade up Mar 20'26_finance confirm.xlsx", data_only=True)
ws = wb['Upgrade BOI10&11 to 15']

# Map detail process names -> Summary machine types
process_map = {
    'SAW': 'SAW', 'SAW QFN': 'SAW', 'SAW(MMT)': 'SAW',
    'DIE ATTACH': 'DIEBONDER', 'DIE ATTACH (MMT)': 'DIEBONDER',
    'WIRE BONDING': 'Wirebonder', 'WIRE BONDING (MMT)': 'Wirebonder',
    'Module (SPG)': 'Module - Reflow',
    'Final leak (SPG)': 'Module - Reflow',
    'Backgrind': 'BACKGRIND',
    'Mounter': 'WAFERMOUNT',
    'MOLD': 'MOLD',
    'Plating': 'SOLDERPLATE',
    'MARK': 'LASERMARK',
    'TRIM/FORM': 'TRIMFORM',
    'Form/Sing': 'TRIMFORM',
    'ISOLATE': 'ISOLATE',
}

# Additional machine types from Summary that don't directly map
# DEFLASH, DIE SHEAR TESTER, PLASMA, PMCOVEN, QA, XRAY
# These are sub-items under other processes in the detail sheet
# Let me identify them by looking at detail content

data = defaultdict(lambda: {'plan': 0, 'spent': 0, 'remain': 0, 'count': 0, 'pr_count': 0, 'items': []})
current_process = ''

for r in range(7, 79):
    b = ws.cell(r, 2).value
    if not b:
        continue
    a = ws.cell(r, 1).value
    if a and isinstance(a, str):
        current_process = a.strip()

    plant = ws.cell(r, 14).value or ''
    total_usd = ws.cell(r, 12).value or 0
    spent = float(ws.cell(r, 18).value or 0)
    remain = float(ws.cell(r, 19).value or 0) if ws.cell(r, 19).value else 0
    pr = ws.cell(r, 17).value
    detail = str(b).strip()

    # Refine machine type based on content
    mch = process_map.get(current_process, current_process)

    # Check if item is really XRAY, PLASMA, QA, etc. based on content
    detail_lower = detail.lower()
    if 'x-ray' in detail_lower or 'xray' in detail_lower or 'x ray' in detail_lower:
        mch = 'XRAY'
    elif 'plasma' in detail_lower or 'rf generator' in detail_lower:
        mch = 'PLASMA'
    elif 'protocol 3' in detail_lower or 'pmcoven' in detail_lower or 'pmc' in current_process.lower():
        if 'OVEN' in str(ws.cell(r, 6).value or '') or 'PMC' in str(ws.cell(r, 6).value or ''):
            mch = 'PMCOVEN'
    elif 'die shear' in detail_lower or 'bond shear' in detail_lower or 'wire pull' in detail_lower:
        mch = 'DIE SHEAR TESTER'
    elif 'deflash' in detail_lower:
        mch = 'DEFLASH'
    elif 'microscope' in detail_lower and current_process in ('WIRE BONDING (MMT)', 'WIRE BONDING'):
        mch = 'QA'

    key = (plant, mch)
    data[key]['plan'] += total_usd
    data[key]['spent'] += spent
    data[key]['remain'] += remain
    data[key]['count'] += 1
    if pr:
        data[key]['pr_count'] += 1
    data[key]['items'].append({
        'row': r, 'detail': detail[:60], 'usd': total_usd, 'spent': spent, 'pr': pr
    })

# Print grouped by plant
for plant in ['MTAI', 'MMT']:
    print(f'\n{"="*80}')
    print(f'  {plant} PLANT')
    print(f'{"="*80}')
    items = [(k, v) for k, v in data.items() if k[0] == plant]
    items.sort(key=lambda x: x[0][1])  # sort by machine type
    plant_plan = 0
    plant_spent = 0
    plant_count = 0
    for (p, mch), v in items:
        print(f'\n  {mch:20s} | Items: {v["count"]:2d} | Plan: ${v["plan"]:>12,.2f} | Spent: ${v["spent"]:>10,.2f} | PRs: {v["pr_count"]}')
        for it in v['items']:
            pr_str = f'PR:{it["pr"]}' if it['pr'] else ''
            sp_str = f'Spent:${it["spent"]:,.2f}' if it['spent'] > 0 else ''
            print(f'    - {it["detail"]:60s} ${it["usd"]:>10,.2f}  {pr_str}  {sp_str}')
        plant_plan += v['plan']
        plant_spent += v['spent']
        plant_count += v['count']
    print(f'\n  {"TOTAL":20s} | Items: {plant_count:2d} | Plan: ${plant_plan:>12,.2f} | Spent: ${plant_spent:>10,.2f}')
